from openpyxl import load_workbook
from pathlib import Path

path = Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb = load_workbook(path)
ws = wb['Submission Log']
row = ['Go2Canada', 'https://www.go2canada.com/submit-content-form', 'Direct blog-post form with hCaptcha', 'Rejected or blocked', 'https://www.go2canada.com/blog-public/submit-blog', 'Manual hCaptcha was verified, but the form returned a validation error stating the story must be under 2,000 characters and then reset to an error/400 state despite shorter content. No receipt and no backlink publication.', 'Do not retry in this cycle; revisit only if the form behavior changes', '2026-09-23']
if tuple(row) not in {tuple(r) for r in ws.iter_rows(min_row=2, values_only=True)}:
    ws.append(row)
wb.save(path)
print(path)
