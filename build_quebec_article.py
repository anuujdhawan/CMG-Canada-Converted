from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "quebec_newcomer_hiring_article.docx"


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)
    return p


def add_source(doc, label, url, checked):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(label + " — ")
    add_hyperlink(p, url, url)
    p.add_run(f" (checked {checked})")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.12

for name, size, bold in [("Title", 22, True), ("Heading 1", 15, True), ("Heading 2", 12, True)]:
    st = styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("What Quebec Small Businesses Should Know Before Hiring Newcomer Founders and Professionals")
remove_paragraph_borders(title)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle.add_run("A practical guide to clearer recruitment, credential conversations, and inclusive onboarding").italic = True

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(14)
meta.add_run("By Anooj Dhawan\n").bold = True
meta.add_run("Content and outreach professional working with Commonwealth Migration Group Inc.")

add_body(doc, "For a small business in Quebec, hiring a newcomer founder or internationally trained professional is often less about finding a perfect résumé and more about asking better questions. Immigration status, language, education, work history, and entrepreneurial experience are related parts of a person’s story, but they are not interchangeable. A fair process separates each issue, checks what the role actually requires, and gives the candidate a clear path to demonstrate capability.")

doc.add_paragraph("Start with the role, not assumptions", style="Heading 1")
add_body(doc, "A job description should explain the work a person will do, the outcomes expected in the first months, the skills that are genuinely essential, and the conditions of employment. Avoid using immigration status as a shortcut for judging reliability, seniority, or suitability. A person who founded a company abroad may bring commercial judgment, supplier relationships, customer insight, or leadership experience even if the Canadian title for that work is unfamiliar.")
add_body(doc, "Write requirements in observable terms. Instead of asking for “Canadian experience” without explaining why, identify the knowledge or task behind the phrase: familiarity with a regulated process, experience with a particular software system, ability to communicate with a defined customer group, or knowledge of Quebec workplace requirements. This makes the process more transparent and gives candidates a meaningful opportunity to explain transferable experience.")

doc.add_paragraph("Separate immigration questions from hiring questions", style="Heading 1")
add_body(doc, "A candidate’s right to work is a legal eligibility question, while their ability to perform a role is a recruitment question. They should be handled separately and consistently. Employers should ask only for the information needed to verify work authorization and should avoid promising immigration outcomes. The Government of Quebec explains that the rules and authorizations for hiring foreign workers depend on the situation, and employers should complete the required steps before work begins.")
add_body(doc, "For roles involving a founder, owner, or professional who is moving to Quebec, there may also be separate business or immigration processes. A person may be building a business, seeking employment, or doing both at different times. Do not treat a business plan, a work permit, a professional licence, and an employment agreement as substitutes for one another. When a question is about immigration status or authorization, direct the person to current government information or an authorized professional rather than improvising an answer.")

doc.add_paragraph("Assess transferable skills with evidence", style="Heading 1")
add_body(doc, "International experience becomes easier to assess when the conversation focuses on evidence. Ask the candidate to describe a project, the problem they faced, the decisions they made, and the result. For a founder, useful prompts may include how they acquired customers, managed cash flow, handled vendors, supervised staff, or adapted a product to a new market. For a professional, ask about the scope of practice, tools used, standards followed, and situations in which they had to make a careful judgment.")
add_body(doc, "Credentials still matter, but the way they matter depends on the occupation. Some jobs are regulated and require authorization from a provincial or territorial regulator before a person can practise. Other jobs are not regulated and may allow an employer to assess a candidate directly. The Government of Canada’s foreign credential recognition resources and Job Bank tools can help employers and candidates understand whether an occupation is regulated and what additional steps may apply.")

doc.add_paragraph("Make language expectations specific and fair", style="Heading 1")
add_body(doc, "Quebec workplaces operate in a language environment where French may be essential to the role, the customers, or the legal context. Explain the language level and tasks required instead of using vague labels such as “native speaker.” Distinguish between customer-facing communication, technical writing, internal collaboration, and language that can be learned through training. If the role has a genuine French requirement, state it early and apply it consistently to every candidate.")
add_body(doc, "Language development can be part of onboarding. A small business might provide written procedures, pair a new employee with a colleague, use plain-language instructions, or identify a few high-priority terms for the first month. These steps do not reduce standards; they help the employee demonstrate the standard more reliably.")

doc.add_paragraph("Handle documents responsibly", style="Heading 1")
add_body(doc, "Recruitment often involves passports, permits, diplomas, reference letters, and financial or business records. Ask for the minimum documentation needed at each stage. Explain who will see it, why it is being collected, how it will be stored, and when it will be deleted or returned. Do not ask for original documents when a secure copy or a verification service is sufficient, and do not retain personal records simply because they may be useful later.")
add_body(doc, "The same principle applies to business founders. A request for a business plan or financial information should be linked to a clear purpose such as evaluating a partnership, a role, or a commercial proposal. Keep immigration documents out of a general hiring file unless they are needed for a lawful work-authorization check.")

doc.add_paragraph("Build an onboarding plan that supports retention", style="Heading 1")
add_body(doc, "A good onboarding plan turns a hiring decision into a workable first ninety days. Clarify the reporting line, priorities, decision-making authority, working language, meeting norms, and how feedback will be given. For an internationally trained professional, identify any role-specific training, supervised practice, or licensing boundary. For a founder joining a small team, agree in writing how strategic input, operational responsibilities, and authority will be divided.")
add_body(doc, "Schedule short check-ins at the end of the first week, first month, and third month. Ask what is clear, what remains unclear, and what support would remove a practical barrier. This is especially useful when a newcomer is learning a new workplace culture or professional vocabulary. It also gives the employer an early opportunity to correct an unclear process instead of treating a preventable misunderstanding as a performance problem.")

doc.add_paragraph("A short checklist for Quebec employers", style="Heading 1")
checklist = [
    "Define the role and essential skills in observable terms.",
    "Separate work authorization from assessment of ability.",
    "Check whether the occupation is regulated and identify the relevant regulator.",
    "Assess international experience through examples and outcomes.",
    "State French and other language expectations clearly and consistently.",
    "Request only the documents needed for the current decision.",
    "Put responsibilities, reporting lines, and onboarding support in writing.",
    "Review current federal and Quebec government guidance before acting on an immigration question.",
]
for item in checklist:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(4)
    p.add_run("• ").bold = True
    p.add_run(item)

doc.add_paragraph("Conclusion", style="Heading 1")
add_body(doc, "Small businesses do not need to become immigration advisers to recruit responsibly. They do need a clear role, a consistent process, careful handling of documents, and the humility to verify questions that fall outside ordinary hiring practice. When employers focus on evidence rather than assumptions, newcomer founders and internationally trained professionals have a fairer opportunity to show what they can contribute—and businesses gain a more reliable way to make decisions.")

doc.add_paragraph("Official sources checked on 28 September 2026", style="Heading 1")
add_source(doc, "Government of Canada Job Bank, Recruit newcomers to Canada", "https://www.jobbank.gc.ca/hiring/newcomers", "28 September 2026")
add_source(doc, "Government of Quebec, Hire an immigrant or a person from a visible minority", "https://www.quebec.ca/entreprises-et-travailleurs-autonomes/administrer-gerer/embauche-gestion-personnel/recruter/embaucher-immigrant", "28 September 2026")
add_source(doc, "Government of Quebec, Settle and integrate in Quebec", "https://www.quebec.ca/en/immigration/settle-and-integrate-in-quebec", "28 September 2026")
add_source(doc, "Government of Canada, Find a job in Canada as a newcomer", "https://www.jobbank.gc.ca/trouverunemploi/nouveaux-arrivants", "28 September 2026")

doc.add_paragraph("Author bio", style="Heading 1")
add_body(doc, "Anooj Dhawan is a content and outreach professional working with Commonwealth Migration Group Inc. He supports communications and editorial partnerships relating to newcomer and relocation information. His work focuses on practical, source-led writing for people making work, business, and settlement decisions in Canada.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Editorial draft | Commonwealth Migration Group Inc.").font.size = Pt(8)

doc.save(OUT)
print(OUT)
