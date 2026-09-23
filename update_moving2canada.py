from openpyxl import load_workbook
from datetime import datetime
p='/Users/themacintosh/Documents/My Files/Projects/CMG-Tailwind-components(V2)/outputs/backlink-outreach-progress-2026-09-23.xlsx'
wb=load_workbook(p); ws=wb['Prospects']; headers=[c.value for c in ws[1]]; headers[0]='Website name'; idx={h:i+1 for i,h in enumerate(headers)}
target=None
for r in range(2,ws.max_row+1):
    if ws.cell(r,idx['Website name']).value=='Moving2Canada': target=r; break
if target is None: raise SystemExit('row not found')
ws.cell(target,idx['Link status']).value='Blocked; email address protected'
ws.cell(target,idx['Notes']).value=('Current page invites original 600–800 word Canada immigration/employment/life contributions and allows 1–2 bio links. '
 'The listed email is Cloudflare-protected in the public page and resolved only to a redacted placeholder; no address was guessed and no message was sent.')
ws.cell(target,idx['Playwright readiness']).value='Inspected 2026-09-23; email route blocked by protected address'
wb['Submission Log'].append(['Moving2Canada','https://moving2canada.com/about-us/become-our-partner/share-your-experience/','Email-only; editorial review','Blocked; not submitted','https://moving2canada.com/about-us/become-our-partner/share-your-experience/','Current rules are viable, but the listed email is Cloudflare-protected and not recoverable as a usable address from the rendered page. No email sent.','Do not guess the address; revisit only if a public contact route becomes available.',datetime(2026,9,23)])
wb.save(p); print('updated',target)
