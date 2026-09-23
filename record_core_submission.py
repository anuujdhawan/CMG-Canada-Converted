from openpyxl import load_workbook
from pathlib import Path

details = Path('outputs/backlink-created-details.xlsx')
wb = load_workbook(details)
ws = wb['Backlink Details']
row = ['https://coremagazines.com/submissions/', 'https://commonwealthmigration.ca/immigrate/express-entry']
if row not in [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]:
    ws.append(row)
wb.save(details)

progress = Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb2 = load_workbook(progress)
log = wb2['Submission Log']
logrow = [
    'Core Magazines',
    'https://coremagazines.com/submissions/',
    'Direct editorial submission form',
    'Submitted/pending',
    'https://coremagazines.com/submissions/?contact-form-sent=34474',
    'Form returned “Thank you for your response” and displayed the submitted name, email, article, and share link. No public article permalink or backlink publication has been verified.',
    'Monitor for publication; do not count as published until publicly verified',
    '2026-09-23',
]
if tuple(logrow) not in {tuple(r) for r in log.iter_rows(min_row=2, values_only=True)}:
    log.append(logrow)
wb2.save(progress)
print(details)
print(progress)
