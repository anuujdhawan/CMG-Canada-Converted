from openpyxl import load_workbook
from openpyxl.styles import Font

p='/Users/themacintosh/Documents/My Files/Projects/CMG-Tailwind-components(V2)/outputs/backlink-outreach-progress-2026-09-23.xlsx'
wb=load_workbook(p)
name='Backlink Submission URLs'
if name in wb.sheetnames:
    ws=wb[name]
else:
    ws=wb.create_sheet(name)
    ws.append(['Website URL','Company backlink URL'])
    for c in ws[1]: c.font=Font(bold=True)
    ws.freeze_panes='A2'
    ws.column_dimensions['A'].width=60
    ws.column_dimensions['B'].width=60
wb.save(p)
print(name, ws.max_row, ws.max_column)
