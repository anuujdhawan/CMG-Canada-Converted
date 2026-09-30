from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = 'newcomer_information_gap_article.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos'); styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(7); styles['Normal'].paragraph_format.line_spacing = 1.12
for name, size in [('Title', 22), ('Heading 1', 15), ('Heading 2', 11.5)]:
    st = styles[name]; st.font.name='Aptos Display' if name=='Title' else 'Aptos'; st._element.rPr.rFonts.set(qn('w:eastAsia'), st.font.name); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(13 if name!='Title' else 0); st.paragraph_format.space_after = Pt(5)

def hyperlink(p, text, url):
    part = p.part; rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr'); color = OxmlElement('w:color'); color.set(qn('w:val'), '1F5A7A'); rPr.append(color); u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u); r.append(rPr)
    t = OxmlElement('w:t'); t.text = text; r.append(t); link.append(r); p._p.append(link)

def add_source(label, url):
    p = doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.18)
    p.add_run('• ').bold=True; hyperlink(p, label, url)

title = doc.add_paragraph(style='Title'); title.add_run('The Newcomer Information Gap')
sub = doc.add_paragraph(); sub.paragraph_format.space_after=Pt(3); r=sub.add_run('How to verify immigration guidance before acting'); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor(70,70,70)
meta = doc.add_paragraph(); meta.alignment=WD_ALIGN_PARAGRAPH.LEFT; meta.paragraph_format.space_after=Pt(14); r=meta.add_run('Draft article | Prepared 28 September 2026 | For editorial review'); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(100,100,100)

p=doc.add_paragraph(); p.add_run('When immigration information changes, the cost of acting on outdated or misleading guidance can be high. ').bold=True; p.add_run('A reliable process is therefore more useful than a confident-sounding answer: identify the source, confirm its date and scope, distinguish general information from advice about a particular person, and verify anyone offering paid help.')

doc.add_heading('Why information becomes difficult to trust', level=1)
doc.add_paragraph('Newcomers often encounter information through search results, social media, messaging groups, employers, schools, friends and commercial websites. These channels can be useful starting points, but they do not all have the same authority. A post may accurately describe a rule that has since changed, simplify an exception into a general rule, or repeat advice copied from another source without showing where it came from.')
doc.add_paragraph('Immigration decisions are also fact-specific. A person’s nationality, status, travel history, family circumstances, documents, deadlines and previous applications may all matter. A general explanation can help someone understand a process, but it cannot by itself determine whether a particular application is complete, eligible or likely to succeed.')

doc.add_heading('Start with a source hierarchy', level=1)
doc.add_paragraph('A practical source hierarchy begins with the Government of Canada and the department responsible for the relevant process, Immigration, Refugees and Citizenship Canada (IRCC). Start at Canada.ca, locate the page for the specific application or service, and read the eligibility, document and submission instructions together. Use the page’s publication or update information as a prompt to re-check the material before acting.')
doc.add_paragraph('Next, check the regulator or professional body when the question concerns a representative. IRCC explains that paid immigration and citizenship consultants must be members in good standing of the College of Immigration and Citizenship Consultants (CICC), while lawyers and notaries are regulated through the relevant law society or the Chambre des notaires du Québec. A person’s title, website or social-media profile is not a substitute for an independent register check.')
doc.add_paragraph('Third-party articles can provide context, questions to ask or explanations in plain language. They should not replace the current official instructions. Be cautious with undated pages, screenshots without links, “secret” shortcuts, copied checklists and advice that conflicts with the government’s forms or guidance.')

doc.add_heading('General information is not personalized advice', level=1)
doc.add_paragraph('General information describes a public process in broad terms. Personalized immigration advice applies rules to an individual’s facts and may involve assessing eligibility, identifying risks, selecting a strategy, preparing forms, responding to procedural issues or communicating with authorities. The distinction matters because a reader may mistake an educational article, chatbot response or informal opinion for a professional assessment.')
doc.add_paragraph('A responsible information provider should state what the material does and does not do. It should identify the date checked, link to the underlying official source, avoid presenting uncertain outcomes as certain, and recommend qualified help when the issue is case-specific or high-risk. No representative can guarantee that IRCC will approve an application, and applicants remain responsible for the truth and completeness of information submitted in their name.')

doc.add_heading('Verify credentials and the scope of service', level=1)
doc.add_paragraph('Before paying for assistance, search the adviser’s name in the relevant regulator’s public register and confirm that the status is current and in good standing. For an immigration consultant, use the CICC Public Register. Confirm that the person you speak with is the person listed, not merely an employee, marketer or referral partner. Ask who will perform the work, who will sign or review it, and who is accountable for the service.')
doc.add_paragraph('Ask for a written agreement that explains the services, fees, taxes, anticipated disbursements, communication arrangements, refund terms and limits of the engagement. A transparent provider should be able to explain what is included without pressuring someone to pay immediately. If an organization has multiple offices, brands or related companies, ask which legal entity is contracting with you and which licensed professional, if any, is responsible for regulated work.')

doc.add_heading('Warning signs that deserve a pause', level=1)
for item in [
    'A promise of guaranteed approval, a job, a visa, permanent residence or faster processing.',
    'Pressure to pay immediately, send money to a personal account or use an unusual transfer method.',
    'Advice to hide facts, provide altered documents or submit information that is not true.',
    'Credentials that cannot be independently verified, or a refusal to identify the responsible professional.',
    'Requests for extensive personal documents before the purpose, security and retention arrangements are explained.',
    'A claim that a private connection can obtain special treatment from an immigration officer.'
]:
    doc.add_paragraph(item, style='List Bullet')
doc.add_paragraph('These are not merely matters of style. IRCC warns that applicants are responsible for information in their applications even when a representative completes them, and that false or misleading information can lead to refusal and other serious consequences.')

doc.add_heading('Protect identity and personal documents', level=1)
doc.add_paragraph('Passports, identity documents, financial records, employment letters, language results and family records can expose someone to identity theft or other harm if mishandled. Share only what is necessary for a clearly defined purpose. Confirm the recipient, use a secure transmission method, keep a record of what was sent, and ask how files are stored, who can access them and when they will be deleted or returned. Do not assume that a website is trustworthy merely because it uses professional branding or a padlock icon.')
doc.add_paragraph('If a document is requested, ask why it is needed and whether a redacted copy or alternative will meet the purpose. Never allow anyone to change, conceal or fabricate information. If something appears fraudulent, preserve relevant messages and receipts and use the reporting channels identified by IRCC and the appropriate regulator.')

doc.add_heading('A pre action checklist', level=1)
table = doc.add_table(rows=1, cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
hdr=table.rows[0].cells; hdr[0].text='Before you act'; hdr[1].text='What to confirm'
rows=[('Source','Is the information on an official, current page? What date was it checked?'),('Fit','Is this general information, or does it require analysis of my personal facts?'),('Representative','Is the person authorized and in good standing in the relevant public register?'),('Agreement','Do I have a written scope of work, fee schedule and responsible contact?'),('Documents','Why is each document needed, and how will it be protected?'),('Truthfulness','Have I reviewed every answer and document for accuracy before submission?'),('Red flags','Has anyone promised an outcome, special access or faster processing?')]
for a,b in rows:
    cells=table.add_row().cells; cells[0].text=a; cells[1].text=b
for row in table.rows:
    for i,c in enumerate(row.cells):
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs: p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
        if row==table.rows[0]:
            tcPr=c._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'1F5A7A'); tcPr.append(shd)
            for p in c.paragraphs:
                for rr in p.runs: rr.font.bold=True; rr.font.color.rgb=RGBColor(255,255,255)

doc.add_heading('A responsible standard for relocation information', level=1)
doc.add_paragraph('Good relocation information does not need to sound certain to be useful. It should make its evidence visible, explain its limits and help a reader decide what to verify next. For newcomers, that means using official sources as the foundation, treating third-party guidance as context, checking credentials independently and pausing when a claim sounds too good to be true.')
doc.add_paragraph('For relocation professionals, the same standard builds trust. Clear boundaries between education, administrative support and regulated professional advice protect readers and providers alike. The goal is not to eliminate every uncertainty; it is to prevent avoidable uncertainty from becoming an expensive or irreversible decision.')

doc.add_heading('Official sources checked 28 September 2026', level=1)
add_source('IRCC Using an immigration and citizenship representative', 'https://www.canada.ca/en/immigration-refugees-citizenship/services/immigration-citizenship-representative.html')
add_source('IRCC Find out if your representative is authorized', 'https://www.canada.ca/en/immigration-refugees-citizenship/services/immigration-citizenship-representative/choose/authorized.html')
add_source('IRCC Online and telephone immigration scams', 'https://www.canada.ca/en/immigration-refugees-citizenship/services/protect-fraud/internet-email-telephone.html')
add_source('IRCC Consequences of immigration and citizenship fraud', 'https://www.canada.ca/en/immigration-refugees-citizenship/services/protect-fraud/consequences-fraud.html')
add_source('IRCC How to report scams fraud or abuse', 'https://www.canada.ca/en/immigration-refugees-citizenship/services/protect-fraud/report-fraud.html')
add_source('CICC Public Register', 'https://register.college-ic.ca/Public-Register-EN')
add_source('Office of the Privacy Commissioner of Canada Identity Theft and You', 'https://www.priv.gc.ca/media/2034/guide_idt_e.pdf')

doc.add_heading('Author information', level=1)
doc.add_paragraph('Anooj Dhawan works with the content and outreach team at Commonwealth Migration Group Inc., operating in Canada as Commonwealth Migration Canada. He is not an RCIC. The organization’s Canadian practice is CICC-regulated, with licensed consultant Pankaj Khanna, RCIC R711256. The following reference link is offered for editorial review and inclusion only if considered relevant:')
p=doc.add_paragraph(); hyperlink(p, 'Commonwealth Migration about page', 'https://commonwealthmigration.ca/about/about-commonwealth-migration')
doc.add_paragraph('Editorial note: This article is educational and does not provide personalized immigration advice, recommend a pathway or promise an outcome. It should be reviewed for publication, links and credentials by the receiving editor.')

doc.save(OUT)
print(OUT)
