from openpyxl import load_workbook
from pathlib import Path

path = Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb = load_workbook(path)
ws = wb['Submission Log']
rows = [
    ['Business Magazine 24', 'https://businessmagazine24.com/submit-article/', 'Not suitable: explicit AI-content restriction', 'Not submitted', 'https://businessmagazine24.com/submit-article/', 'The current guidelines reject raw LLM output and require the writing, judgment, and on-the-record sourcing to be the contributor’s own; no submission made.', 'Exclude from current campaign', '2026-09-23'],
    ['The Atlantic Current', 'https://theatlanticcurrent.com/submit-content-new/', 'Blocked before form', 'Not submitted', 'https://theatlanticcurrent.com/submit-content-new/', 'Direct form includes an image CAPTCHA and is intended for local Atlantic Canada content; no submission made because the route is not a fit for a generic nationwide article.', 'Exclude from current campaign', '2026-09-23'],
    ['The Hamilton Post Magazine', 'https://www.thehamiltonpostmagazine.com/submit', 'Not suitable: local editorial fit', 'Not submitted', 'https://www.thehamiltonpostmagazine.com/submit', 'Form accepts 500–750 word local or lifestyle submissions, but the stated categories and local focus do not establish a genuine Hamilton connection for this campaign.', 'Exclude from current campaign', '2026-09-23'],
]
existing = {tuple(r) for r in ws.iter_rows(min_row=2, values_only=True)}
for row in rows:
    if tuple(row) not in existing:
        ws.append(row)
wb.save(path)
print(path)
