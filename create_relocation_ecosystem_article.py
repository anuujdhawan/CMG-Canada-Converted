from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path("outputs/relocation-ecosystem-information-infrastructure-article.docx")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_source(doc, label, url, checked):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{label} — checked {checked}: ")
    add_hyperlink(p, url, url)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(8)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name in ["Title", "Heading 1", "Heading 2"]:
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name].font.color.rgb = None

title = doc.add_paragraph(style="Title")
title.add_run("The Information Handoff Gap in International Relocation")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("Why source verification and clear boundaries matter before people act on migration information")
run.italic = True

doc.add_paragraph(
    "People planning an international move rarely experience policy as a single document. "
    "They encounter a chain of explanations: a government page, a search result, a community post, "
    "an employer’s advice, and sometimes a professional consultation. Each handoff can add useful context, "
    "but each can also introduce ambiguity. That information handoff gap is a policy problem because a small "
    "misunderstanding can affect money, timing, employment, and a family’s decision to move."
)

doc.add_heading("Information is part of the relocation system", level=1)
doc.add_paragraph(
    "Immigration and relocation policies are often evaluated through outcomes such as admission numbers, "
    "labour-market participation, or settlement indicators. The communication layer that connects people to "
    "those policies receives less attention. Yet a rule that cannot be found, understood, or applied to the "
    "right question is difficult to use in practice. This matters internationally: policy information is read "
    "in many countries, translated through informal networks, and interpreted by employers, schools, community "
    "organizations, and private publishers."
)
doc.add_paragraph(
    "The official source remains essential, but it is not always sufficient on its own. Government pages may be "
    "accurate while still assuming that readers know the difference between an eligibility requirement, a "
    "program description, and an individual assessment. People comparing destinations from abroad may also find "
    "several accurate pages without an obvious explanation of how the pages relate to one another."
)

doc.add_heading("The risk grows at each handoff", level=1)
doc.add_paragraph(
    "A search result can omit a condition. A social post can repeat an old rule. A summary can turn a possibility "
    "into a promise. A well-intentioned employer can confuse a work-permit question with a permanent-residence "
    "question. Once an error is repeated, readers may encounter it in several places and mistake repetition for "
    "confirmation."
)
doc.add_paragraph(
    "The consequences are not limited to an incorrect form. People may pay for translations, language tests, "
    "credential assessments, travel, or housing based on an assumption that later proves wrong. Families may delay "
    "a decision or make one too quickly. These are ordinary policy effects experienced through an information "
    "system rather than through a statute or regulation alone."
)

doc.add_heading("Four practices that make relocation information safer", level=1)
doc.add_paragraph(
    "First, every practical explanation should identify its primary source and show when the source was checked. "
    "A visible date does not make a page permanently current, but it gives readers a starting point for judging "
    "whether further verification is needed."
)
doc.add_paragraph(
    "Second, writers should explain the boundary between general information and professional advice. A checklist "
    "can help a reader understand what to research; it cannot determine eligibility, interpret every fact pattern, "
    "or predict an official decision. Clear boundaries reduce both overconfidence and unnecessary fear."
)
doc.add_paragraph(
    "Third, publishers should plan for updates instead of treating corrections as exceptional. A page about a "
    "time-sensitive policy needs an owner, a review habit, and a way for readers to flag stale information. "
    "The source hierarchy should be visible: official requirements first, independent explanation second, and "
    "an invitation to seek individualized help where the facts are complex."
)
doc.add_paragraph(
    "Fourth, explanations should be organized around decisions rather than labels alone. Readers need to know "
    "which question a source answers, what evidence they need before relying on it, and what the source does not "
    "establish. This approach is more useful than collecting isolated program names or promising timelines."
)

doc.add_heading("A shared responsibility across borders", level=1)
doc.add_paragraph(
    "Settlement organizations, journalists, researchers, and professional publishers each have a different role. "
    "Settlement organizations can make official information more accessible without replacing its authority. "
    "Journalists can report the effect of a policy while linking to the underlying document. Researchers can show "
    "where terminology or administrative design creates unequal access to information. Professional publishers "
    "can provide navigation tools if they disclose their interests, avoid guaranteed outcomes, and distinguish "
    "education from individualized advice."
)
doc.add_paragraph(
    "This is also a question of trust. A country’s immigration system can be carefully designed and still appear "
    "arbitrary when people encounter contradictory explanations. Clear sourcing, visible uncertainty, and prompt "
    "corrections make policy more legible. They also make it easier for readers to recognize predatory promises "
    "and to ask better questions before spending money or making an irreversible decision."
)

doc.add_heading("The practical test", level=1)
doc.add_paragraph(
    "Before acting on relocation information, a reader should be able to answer five questions: Who is responsible "
    "for the rule? When was the source last checked? What decision does it actually address? Which facts could change "
    "the answer? Where can the reader obtain individualized advice if needed? A publication that helps readers ask "
    "those questions contributes to better policy outcomes without pretending to replace an official decision-maker."
)
doc.add_paragraph(
    "International relocation will continue to involve difficult debates about mobility, labour, identity, and "
    "belonging. Those debates are mediated by the information people can use. Treating that information as part of "
    "the policy infrastructure—not as an afterthought—would make relocation systems more transparent, more humane, "
    "and easier to evaluate."
)

doc.add_heading("Sources checked", level=1)
add_source(doc, "Immigration Refugees and Citizenship Canada, Express Entry", "https://www.canada.ca/en/immigration-refugees-citizenship/services/immigrate-canada/express-entry.html", "23 September 2026")
add_source(doc, "Government of Canada, Immigration and citizenship", "https://www.canada.ca/en/services/immigration-citizenship.html", "23 September 2026")
add_source(doc, "Settlement.Org, newcomer information and resources", "https://settlement.org/", "23 September 2026")
add_source(doc, "Commonwealth Migration Group, general Express Entry information", "https://commonwealthmigration.ca/immigrate/express-entry", "23 September 2026")

doc.add_heading("Author information", level=1)
doc.add_paragraph(
    "Anooj Dhawan works in content and outreach for Commonwealth Migration Group Inc. His role is editorial and "
    "communications-focused; he is not presenting himself as a licensed immigration consultant. The organizational "
    "relationship and any reference link are subject to independent verification and editorial review."
)

doc.save(OUT)
print(OUT)
