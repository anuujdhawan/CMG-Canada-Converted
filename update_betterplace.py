from openpyxl import load_workbook
from datetime import datetime

p = '/Users/themacintosh/Documents/My Files/Projects/CMG-Tailwind-components(V2)/outputs/backlink-outreach-progress-2026-09-23.xlsx'
wb = load_workbook(p)
ws = wb['Prospects']
headers = [c.value for c in ws[1]]
headers[0] = 'Website name'
idx = {h:i+1 for i,h in enumerate(headers)}
target = None
for r in range(2, ws.max_row+1):
    if ws.cell(r, idx['Website name']).value == 'BetterPlace Immigration':
        target = r; break
if target is None:
    raise SystemExit('BetterPlace Immigration row not found')
ws.cell(target, idx['Link status']).value = 'Submitted/pending'
ws.cell(target, idx['Notes']).value = ('Current editorial route accepts expert/feature contributions by email. '
    'Sent a tailored pitch from authorized seolaunchers@gmail.com for “The Newcomer Information Gap: '
    'How to Verify Canadian Immigration Guidance Before Acting,” with one contextual Express Entry reference. '
    'Gmail displayed Message sent; no public article URL or backlink yet.')
ws.cell(target, idx['Playwright readiness']).value = 'Submitted 2026-09-23 via authorized Gmail; pending editorial response'
log = wb['Submission Log']
log.append([
    'BetterPlace Immigration',
    'mailto:info@betterplaceimmigration.com',
    'Email-only; editorial review',
    'Submitted/pending',
    'Gmail Message sent confirmation; no public article URL',
    'Tailored editorial pitch sent from authorized seolaunchers@gmail.com. No public backlink exists yet.',
    'Await editorial response; verify only if a public article is published and the CMG link opens.',
    datetime(2026,9,23)
])
wb.save(p)
print('updated', target, 'log_row', log.max_row)
