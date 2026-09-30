from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = '/Users/themacintosh/Documents/My Files/Projects/CMG-Tailwind-components(V2)/information-handoff-gap-revised.docx'
CHECKED='23 September 2026'

def hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1F4E79'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    r.append(rPr); t = OxmlElement('w:t'); t.text = text; r.append(t); link.append(r)
    paragraph._p.append(link)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(item)

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(7); styles['Normal'].paragraph_format.line_spacing=1.08
for s in ['Title','Heading 1','Heading 2']:
    styles[s].font.name='Aptos'; styles[s].font.color.rgb=RGBColor(0,0,0)
styles['Title'].font.size=Pt(21); styles['Title'].font.bold=True
styles['Heading 1'].font.size=Pt(14); styles['Heading 1'].font.bold=True; styles['Heading 1'].paragraph_format.space_before=Pt(13); styles['Heading 1'].paragraph_format.space_after=Pt(5)
styles['Heading 2'].font.size=Pt(11.5); styles['Heading 2'].font.bold=True; styles['Heading 2'].paragraph_format.space_before=Pt(9); styles['Heading 2'].paragraph_format.space_after=Pt(3)
title_style_ppr = styles['Title']._element.get_or_add_pPr()
for border in list(title_style_ppr.findall(qn('w:pBdr'))):
    title_style_ppr.remove(border)

p=doc.add_paragraph(style='Title'); p.add_run('The Information Handoff Gap in International Relocation')
# Remove the built-in Word Title style's decorative bottom border.
title_ppr = p._p.get_or_add_pPr()
for border in list(title_ppr.findall(qn('w:pBdr'))):
    title_ppr.remove(border)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; r=p.add_run('How coordinated, source-checked information helps people make informed decisions'); r.italic=True; r.font.size=Pt(11)

doc.add_paragraph('People planning an international move rarely experience policy as a single document. They encounter a chain of explanations: a government page, an employer’s advice, a housing provider’s requirement, a transport instruction, and sometimes a professional consultation. Each handoff can add useful context, but each can also introduce ambiguity. The practical question is not only whether each statement sounds plausible. It is whether the right person verified it, whether its scope is clear, and whether someone owns the next step.')

doc.add_heading('Information is part of the relocation system', level=1)
doc.add_paragraph('Immigration and relocation policies are often evaluated through outcomes such as admission numbers, labour-market participation, or settlement indicators. The communication layer that connects people to those policies receives less attention. Yet a rule that cannot be found, understood, or applied to the right question is difficult to use in practice. This matters internationally: policy information is read in many countries, translated through informal networks, and interpreted by employers, schools, community organizations, housing providers, transport companies, and private publishers.')
doc.add_paragraph('The official source remains essential, but it is not always sufficient on its own. Government pages may be accurate while still assuming that readers know the difference between an eligibility requirement, a program description, a document requirement, and an individual assessment. A coordinator’s job is not to replace the source. It is to help the team keep the source, its date, its scope, and its owner visible.')

doc.add_heading('A hypothetical move across four providers', level=1)
doc.add_paragraph('The following example is fictional and is included to show an information flow, not to provide immigration or relocation advice. It uses a Canadian immigration context only where stated; other jurisdictions may use different rules, authorities, and professional titles.')
rows=[
('Immigration','A licensed Canadian immigration professional identifies the official IRCC page that answers the person’s work-authorisation question and records the page title, URL, date checked, and facts still needing confirmation.','Immigration professional verifies the immigration claim against IRCC.','Immigration professional owns the next action: confirm the person’s facts and tell the employer what remains unverified.'),
('Employment','The employer confirms the job title, location, start date, and any employer-side requirement. It does not treat a general immigration explanation as proof that the person may work.','Employer verifies its own job and onboarding information; immigration questions remain with the immigration professional or official source.','Employer owns the next action: provide the written job details and flag any change immediately.'),
('Housing','A housing provider explains its application documents, deposit terms, move-in date, and identity checks. It does not request a full immigration file when a narrower document will answer the housing question.','Housing provider verifies its rental requirements and the coordinator checks that they match the person’s consent and situation.','Housing provider owns the next action: state what is missing and where the applicant can verify the requirement.'),
('Transport','A transport provider confirms booking rules, baggage limits, delivery dates, and identification needed for the service. These instructions do not establish immigration eligibility.','Transport provider verifies current service rules and the coordinator records the version or date.','Transport provider owns the next action: confirm the booking or explain the change pathway if timing shifts.'),
]
t=doc.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
headers=['Provider','Instruction passed on','Who verifies it','Who owns the next step']
for i,h in enumerate(headers):
    c=t.rows[0].cells[i]; c.text=h; set_cell_shading(c,'1F4E79'); set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(9)
for n,row in enumerate(rows):
    cells=t.add_row().cells
    for i,val in enumerate(row):
        cells[i].text=val; set_cell_margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if n%2==1: set_cell_shading(cells[i],'F2F5F7')
        for p in cells[i].paragraphs:
            for run in p.runs: run.font.size=Pt(8.5)

doc.add_heading('When instructions conflict or change', level=2)
doc.add_paragraph('Suppose the employer changes the start date after the housing provider has issued a move-in date. The coordinator should not silently choose which instruction “wins.” The coordinator records both versions, identifies the source and date of each, asks the relevant owner to confirm the change, and tells the newcomer what decision is paused. If the conflict touches immigration status, work authorization, admissibility, or another regulated question, the matter returns to the appropriate official source or qualified professional. The record should show the resolution, the person who confirmed it, the date, and the next action.')
doc.add_paragraph('A useful handoff is therefore a small, auditable record: source, date checked, scope, facts assumed, responsible person, next action, and change path. It helps the newcomer see what is known, what is not yet known, and who is accountable for finding out.')

doc.add_heading('Credential verification and sensitive documents', level=1)
doc.add_paragraph('People moving across borders are often asked to trust unfamiliar professionals. In Canada, a person seeking paid immigration advice should independently verify whether the professional is authorized to provide that service. The College of Immigration and Citizenship Consultants maintains a public register for regulated Canadian immigration consultants; lawyers and paralegals are governed through their respective law societies. The appropriate register and professional category can differ by jurisdiction and service. Check the relevant regulator directly rather than relying only on a badge, logo, directory listing, or referral.')
doc.add_paragraph('Readers and coordinators should also avoid sharing more information than the task requires. Before sending a document, ask: What decision is this document needed for? Who is the recipient? What is the minimum page or field needed? How will it be stored, who can access it, and when will it be deleted or returned? Use secure channels where available, redact unrelated identifiers when permitted, and keep a record of consent. A passport scan, financial record, medical record, or immigration file should not be circulated to every provider in the relocation chain merely for convenience.')
doc.add_paragraph('These are general information-management practices, not a substitute for personalized legal, immigration, privacy, housing, employment, or transportation advice. When the answer depends on the person’s detailed facts, the appropriate professional should explain the scope of the advice and the documents genuinely required.')

doc.add_heading('Canadian scope and source boundaries', level=1)
doc.add_paragraph('The Canadian examples in this article are limited to information published by Canadian authorities and the Canadian professional-regulatory context. Immigration rules, credential titles, privacy obligations, housing practices, and transport requirements vary by country. A Canadian source should not be presented as a universal answer, and an explanation from a private organization should not be treated as an official decision.')
doc.add_paragraph('For example, IRCC’s Express Entry page describes the federal system and its requirements, but it does not decide a particular person’s eligibility without the person’s complete facts. The Government of Canada’s immigration and citizenship portal provides the broader official entry point. Direct links and checked dates are listed below so readers can revisit the source when a rule or page changes.')

doc.add_heading('Closing checklist for a safer handoff', level=1)
add_bullets(doc,[
'Source: What is the original source of this instruction, and is it official, professional, provider-specific, or a secondary explanation?',
'Date: When was it checked, and is there a clear way to verify whether it has changed?',
'Scope: Which decision does it address, which jurisdiction does it apply to, and what does it not establish?',
'Responsible person: Who verified it and who owns the next step?',
'Next action: What exactly must happen next, by whom, and by what date?',
'Change path: Where should the newcomer or coordinator go if the instruction conflicts with another one or becomes outdated?',
'Documents: What is the minimum information needed, who has consent to see it, and how will it be protected?'
])
doc.add_paragraph('A publication that helps readers ask these questions contributes to better policy outcomes without pretending to replace an official decision-maker or a personalized professional assessment.')

doc.add_heading('Sources checked', level=1)
sources=[
('Immigration, Refugees and Citizenship Canada, Express Entry', 'https://www.canada.ca/en/immigration-refugees-citizenship/services/immigrate-canada/express-entry.html'),
('Government of Canada, Immigration and citizenship', 'https://www.canada.ca/en/services/immigration-citizenship.html'),
('College of Immigration and Citizenship Consultants, Find an Immigration Consultant', 'https://register.college-ic.ca/'),
('Settlement.Org, newcomer information and resources', 'https://settlement.org/')]
for label,url in sources:
    p=doc.add_paragraph(style='List Bullet'); p.add_run(f'{label} — checked {CHECKED}: '); hyperlink(p,url,url)

doc.add_heading('Author information', level=1)
doc.add_paragraph('Anooj Dhawan works in content and outreach for Commonwealth Migration Group Inc. His role is editorial and communications-focused; he is not presenting himself as a licensed immigration consultant. The organizational relationship and any reference link are provided for independent verification and editorial review.')
p=doc.add_paragraph(); p.add_run('Professional reference for the organization: Pankaj Khanna, RCIC, Licence No. R711256, is identified by Commonwealth Migration Group as a Regulated Canadian Immigration Consultant. Readers should verify current status independently through the '); hyperlink(p,'CICC public register','https://register.college-ic.ca/Public-Register-EN/Licensee/Profile.aspx?ID=14725'); p.add_run('.')
p=doc.add_paragraph(); p.add_run('Optional company reference, separate from the evidence sources: '); hyperlink(p,'Commonwealth Migration Group Inc. — Express Entry information','https://commonwealthmigration.ca/immigrate/express-entry')
doc.add_paragraph(f'Checked {CHECKED}. Publication and link inclusion remain subject to editorial review.')

doc.core_properties.title='The Information Handoff Gap in International Relocation'
doc.core_properties.author='Anooj Dhawan'
doc.core_properties.subject='Information coordination in international relocation'
doc.save(OUT)
print(OUT)
