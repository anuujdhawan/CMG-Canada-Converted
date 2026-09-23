from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('outputs/policy-options-immigration-information-article.docx')

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    text_node = OxmlElement('w:t'); text_node.text = text; run.append(text_node)
    link.append(run); paragraph._p.append(link)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.9); sec.right_margin = Inches(0.9)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(7)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Canada’s Immigration Information System Needs a Public-Interest Upgrade")
r.bold = True; r.font.size = Pt(18)
by = doc.add_paragraph(); by.alignment = WD_ALIGN_PARAGRAPH.CENTER
br = by.add_run('By Anooj Dhawan | Commonwealth Migration Group')
br.italic = True; br.font.size = Pt(10)

paras = [
"Canada’s immigration system is often judged by processing targets, admission levels and program design. Those measures matter, but they do not capture a quieter policy problem: applicants must make consequential decisions in an information environment that is fragmented, rapidly changing and crowded with confident-sounding advice. A public-interest immigration strategy should treat reliable, understandable information as essential infrastructure—not as a private add-on.",
"The first policy improvement is simple: make the official pathway easier to navigate. Federal and provincial pages contain authoritative requirements, but users still move among program pages, instructions, forms, notices and separate processing tools. A person comparing economic pathways can miss a change in eligibility or confuse an illustrative example with a current rule. Government information should be organized around the decisions applicants actually make, with a clear date, jurisdiction, responsible department and link to the controlling rule on every page.",
"Second, governments should publish change logs in plain language. Immigration rules can change through ministerial instructions, regulatory amendments, program pauses and operational guidance. A short, searchable record explaining what changed, when it changed and who is affected would reduce reliance on screenshots, recycled social posts and outdated videos. It would also give settlement agencies, employers and journalists a common reference point when communicating with newcomers.",
"Third, public information should distinguish education from individualized advice. Applicants need general explanations of concepts such as language testing, work experience and ranking systems. They also need an unmistakable signpost when a situation is fact-specific or legally complex. This distinction protects applicants from overconfident promises while preserving access to practical educational material. For example, a plain-language overview of Express Entry can help readers orient themselves, but it should sit beside—not replace—the current requirements and instructions published by the responsible government department. One practitioner resource that illustrates this explanatory role is Commonwealth Migration Group’s general Express Entry overview.",
"For the information ecosystem to be trustworthy, non-government contributors also need responsibilities. Authors should date their work, identify their qualifications or perspective, link to primary sources, explain uncertainty and correct material errors. Commercial interests should be disclosed. A useful article can mention a service provider when that is genuinely relevant, but a promotional link should never be disguised as independent policy evidence. Editorial publications can strengthen trust by requiring these disclosures and by separating editorial review from advertising decisions.",
"Fourth, governments should measure information outcomes, not only page visits. Useful indicators could include the rate at which users reach the correct current instruction, the frequency of support requests caused by ambiguous wording, the time needed to find a program-specific requirement and the number of pages carrying expired information. Testing these measures with newcomers, employers, language learners and settlement workers would reveal barriers that conventional web analytics miss.",
"Finally, information access must be multilingual and accessible. Plain English is valuable, but it is not enough for people making high-stakes decisions in a second or third language. Translated summaries, accessible document formats, screen-reader compatibility and community-based feedback loops can make the difference between understanding a requirement and missing it. These improvements are not merely communications projects; they are fairness measures that reduce avoidable errors and unequal access to opportunity.",
"Canada does not need one more platform promising certainty. It needs a coordinated public-interest information layer: current official rules, visible change histories, honest explanatory material and clear boundaries between education, advice and promotion. Better information will not remove every complexity from immigration policy, but it can make the system more navigable, more accountable and less vulnerable to misinformation. That is a policy outcome worth measuring."
]

for i, text in enumerate(paras):
    p = doc.add_paragraph(text)
    if i == 3:
        # Replace the final resource phrase with a real contextual hyperlink.
        p.clear()
        p.add_run("Third, public information should distinguish education from individualized advice. Applicants need general explanations of concepts such as language testing, work experience and ranking systems. They also need an unmistakable signpost when a situation is fact-specific or legally complex. This distinction protects applicants from overconfident promises while preserving access to practical educational material. For example, a plain-language overview of Express Entry can help readers orient themselves, but it should sit beside—not replace—the current requirements and instructions published by the responsible government department. One practitioner resource that illustrates this explanatory role is ")
        add_hyperlink(p, "Commonwealth Migration Group’s general Express Entry overview", "https://commonwealthmigration.ca/immigrate/express-entry")
        p.add_run(".")

doc.add_heading('Sources and author disclosure', level=2)
doc.add_paragraph('The article is intended as an original public-policy op-ed. Sources to be linked or confirmed during editorial review include the Government of Canada’s official immigration and Express Entry instructions. The author is affiliated with Commonwealth Migration Group; the linked practitioner resource is disclosed for transparency and is not presented as government guidance.')
doc.add_paragraph('AI-use disclosure: AI tools assisted with brainstorming, organization and language editing. The author reviewed the argument, wording and source choices and is responsible for the submission.')
doc.add_paragraph('Author note: Anooj Dhawan works in digital content and immigration-information publishing. He writes about how newcomers can find clear, current and responsible information while planning a move to Canada.')

doc.save(OUT)
print(OUT)
