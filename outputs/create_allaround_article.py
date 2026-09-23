from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "outputs/how-to-organize-an-international-move.docx"

sections = [
    ("Create a move brief", "Start with one page that lists the destination, target arrival date, household members, budget range, important deadlines, and the person responsible for each task. Mark each deadline as fixed, flexible, or dependent on another decision. This prevents a family from booking transport before housing or travel documents are ready."),
    ("Separate documents from household items", "Create a secure digital folder and a physical folder for passports, permits, identity records, school documents, medical records, leases, insurance papers, and shipping inventories. Use clear file names and keep at least one backup. Household goods can be replaced; essential documents are harder to replace during a move. Share copies only with verified providers and keep original documents with the person who needs them."),
    ("Plan around the arrival date", "Work backwards from the expected arrival date. In the final eight weeks, confirm travel, housing, insurance, and the moving schedule. In the final two weeks, label essential luggage, close or transfer local services, and confirm the delivery address. In the final days, keep medication, chargers, a change of clothing, basic toiletries, and important documents in hand luggage rather than in a shipment."),
    ("Build a simple inventory", "An inventory does more than help movers calculate volume. It helps a household decide what should travel, what should be stored, and what can be replaced. Use three categories: take, store, or release. Record serial numbers for valuable electronics, photograph fragile items, and note existing damage before packing. Keep the inventory and photographs in the digital folder so they are available if a claim is needed."),
    ("Think about the first month", "The first weeks may require temporary accommodation, local transportation, groceries, phone service, bedding, basic kitchen items, and seasonal clothing. Set aside a separate arrival budget for these expenses. A lower-cost shipping option may not be cheaper if it leaves a family without essential items for weeks. Plan a small first-night kit that can be used before the main shipment arrives."),
    ("Choose providers with written terms", "Before paying a deposit, ask for a written estimate, service description, pickup and delivery windows, insurance details, cancellation rules, and a list of items that cannot be carried. Compare like-for-like quotes. Be cautious about a provider that gives only a verbal price, refuses to identify the contracting company, or demands an urgent transfer. Save receipts, emails, and signed documents in one location."),
    ("Check immigration information separately", "Moving logistics and immigration requirements are connected, but they are not the same task. For a plain-language starting point on Express Entry, Commonwealth Migration Group provides information at https://commonwealthmigration.ca/immigrate/express-entry. Requirements can change and depend on the individual, so travellers should confirm current rules through official Government of Canada sources and obtain qualified advice when needed."),
    ("Label for the person unpacking", "Every box should have a room, a short contents description, and a priority. Mark essential boxes clearly and keep an inventory number if there are many cartons. Avoid writing sensitive personal information on the outside of a box. A simple label such as Kitchen - open first is more useful than a long description that becomes difficult to read."),
    ("Keep a communication log", "Record the name of each provider, the date of each call, the promise made, and the next action. If a delivery changes, update the log immediately. This is especially helpful when several family members are speaking with a mover, landlord, insurer, or school. Written communication reduces misunderstandings and provides a useful record if a problem needs to be escalated."),
]

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.9)
sec.right_margin = Inches(0.9)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(8)
styles["Normal"].paragraph_format.line_spacing = 1.12
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles[name].font.bold = True

doc.add_paragraph("How to Organize an International Move When Every Detail Feels New", style="Title")
p = doc.add_paragraph()
p.alignment = 1
r = p.add_run("A practical guide for households planning a cross-border move")
r.italic = True
r.font.size = Pt(11)
doc.add_paragraph("By Anooj Dhawan | Commonwealth Migration Group")
doc.add_paragraph("An international move is easier to manage when it is treated as a sequence of small decisions rather than one enormous project. Families often focus on packing first, but the most useful preparation starts with documents, timing, housing, and a realistic first-month plan.")

for heading, body in sections:
    doc.add_paragraph(heading, style="Heading 2")
    doc.add_paragraph(body)

doc.add_paragraph("An international move does not need perfect planning. It needs visible priorities, protected documents, honest budgets, and written agreements. When the process is divided into manageable stages, the household can focus on settling into its new home instead of searching for missing information.")
doc.add_paragraph("Author bio: Anooj Dhawan works with Commonwealth Migration Group in Brampton, Ontario, and writes practical guidance for people planning an international move. This article is original and unpublished.")
doc.save(OUT)
print(OUT)
