from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


OUT = "outputs/attorneys-magazine-tfw-compliance-checklist.docx"


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_mixed(paragraph, pieces):
    for item in pieces:
        if isinstance(item, tuple):
            add_hyperlink(paragraph, item[0], item[1])
        else:
            paragraph.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08

article_title = styles.add_style("ArticleTitle", WD_STYLE_TYPE.PARAGRAPH)
article_title.base_style = styles["Normal"]
article_title.font.name = "Aptos Display"
article_title.font.size = Pt(22)
article_title.font.bold = True
article_title.font.color.rgb = RGBColor(0, 0, 0)
article_title.paragraph_format.space_after = Pt(5)
article_title.paragraph_format.keep_with_next = True

for name in ["Heading 1", "Heading 2"]:
    st = styles[name]
    st.font.name = "Aptos"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(5)

doc.add_paragraph("From work permit to paycheque: a Canadian employer compliance checklist", style="ArticleTitle")
byline = doc.add_paragraph()
run = byline.add_run("By Anuj Dhawan")
run.bold = True
run.font.size = Pt(11)

doc.add_paragraph(
    "An LMIA approval or an employer-specific work permit is not the end of an employer’s responsibilities. It is the point at which the promised job, wage, duties and working conditions must become a real, documented employment relationship. A practical compliance review should therefore follow the worker from the offer letter to the first paycheque—and continue for years afterward."
)

doc.add_paragraph(
    "For employers, the most useful checklist is not a file of forms. It is a short chain of questions: Is the job the same job that was approved? Is the worker being paid what was promised? Are records complete? Has anything changed, and if so, was the change assessed before it took effect?"
)

doc.add_paragraph("1. Start with the approved job, not the current assumption", style="Heading 1")
doc.add_paragraph(
    "The first review should put the LMIA decision letter, any annexes, the employment agreement and the worker’s permit conditions beside the current payroll and job description. The employer should compare the occupation, duties, work location, wage, hours and working conditions line by line. A role can evolve in ordinary business life, but an informal change is not automatically covered by the original approval."
)
doc.add_paragraph(
    "The same discipline applies when a worker moves between supervisors, works at a different location, receives new duties or experiences a change in pay. The question is not simply whether the new arrangement seems reasonable. It is whether the arrangement remains consistent with the program conditions and employment laws that govern the file."
)

doc.add_paragraph("2. Make the paycheque match the promise", style="Heading 1")
doc.add_paragraph(
    "Payroll is often the clearest evidence of what actually happened. Employers should retain wage records, hours, deductions, overtime calculations and proof of payment, and should be able to explain any difference between the approved wage and the amount paid. A worker should not be asked to return wages, pay recruitment costs or absorb an employer’s program-related expenses through an improper deduction."
)
doc.add_paragraph(
    "The federal employer-compliance guidance says inspectors may verify whether the worker received employment in the same occupation and received wages that are mainly the same as, or not less favourable than, the offer. It also requires employers to keep relevant records for six years beginning on the first day of the employment period for which the work permit was issued."
)
p = doc.add_paragraph()
add_mixed(p, [
    "The detailed requirements are set out in ",
    ("Employment and Social Development Canada’s employer-compliance guidance", "https://www.canada.ca/en/employment-social-development/services/foreign-workers/employer-compliance.html"),
    ". The practical lesson is simple: payroll should be reviewed against the approval documents before a discrepancy becomes an inspection question."
])

doc.add_paragraph("3. Treat records as an operating system", style="Heading 1")
doc.add_paragraph(
    "A compliance file should allow a new manager—or an inspector who has never seen the business—to understand what was promised, what was delivered and who approved each change. At minimum, the file should connect the recruitment records, job offer, signed agreement, schedules, time records, payroll, workplace policies, benefit information, leave records and communications about material changes."
)
doc.add_paragraph(
    "This is also a governance issue. Assign an owner for the file, set reminders for document retention and create a change log. If the person who prepared the LMIA leaves the organization, the business should not lose the context needed to prove compliance. A checklist that exists only in one adviser’s inbox is not a durable control."
)

doc.add_paragraph("4. Check the workplace, not just the paperwork", style="Heading 1")
doc.add_paragraph(
    "Compliance is not achieved by matching a document while the workplace tells a different story. Employers should check that the worker has the promised equipment, training, schedule, supervision and safe working conditions. They should make it easy for workers to understand their rights and raise concerns without retaliation. Managers should know that a worker’s immigration status cannot be used to justify threats, withheld wages or a different standard of workplace protection."
)
p = doc.add_paragraph()
add_mixed(p, [
    "The federal overview of ",
    ("temporary foreign worker rights and labour standards", "https://www.canada.ca/en/immigration-refugees-citizenship/services/work-canada/worker-rights.html"),
    " is a useful companion to the employer checklist because it keeps the worker’s perspective visible."
])

doc.add_paragraph("5. Prepare for an inspection before receiving a notice", style="Heading 1")
doc.add_paragraph(
    "An inspection may be triggered by a concern, a history of non-compliance, a random selection or another regulatory reason. It may be announced or unannounced, on-site or virtual, and can include requests for documents and interviews with employers or workers. Waiting for a notice before identifying the records, people and decisions connected to the file creates avoidable risk."
)
doc.add_paragraph(
    "A sensible preparation exercise is a quarterly file walk-through. Ask someone who was not involved in the original application to trace one worker’s record from recruitment to the latest pay period. Can they find the approved duties? Can they reconcile the wage? Can they explain a location or schedule change? Can they show that concerns were addressed? The answers reveal weaknesses while there is still time to correct them."
)

doc.add_paragraph("6. Escalate changes early", style="Heading 1")
doc.add_paragraph(
    "The highest-risk moment is often a change that looks too ordinary to report: a new worksite, a different occupation, a revised wage, a change in hours, a business reorganization or a worker who leaves earlier than expected. The employer should pause, record the facts and obtain advice on whether the change requires a new application, notification, voluntary disclosure or another step."
)
doc.add_paragraph(
    "The employer-compliance regime is designed around evidence. When a business discovers a possible problem, keeping quiet does not make the record cleaner. A documented internal review, prompt correction and appropriate disclosure may give the employer a stronger position than an unexplained gap discovered later."
)

doc.add_paragraph("A checklist is a culture, not a closing task", style="Heading 1")
doc.add_paragraph(
    "The best compliance checklist is used before hiring, during onboarding, when payroll changes, whenever duties or locations shift and at regular review points afterward. It protects the employer, but it also protects the worker whose livelihood depends on the promised job being genuine in practice."
)
doc.add_paragraph(
    "Employers looking for a plain-language overview can also consult Commonwealth Migration Group’s related explanation of TFWP employer compliance; any publication link should be used only if it fits the outlet’s editorial policy and the reader’s need. The central rule remains independent of any service provider: compare the approved terms with the real workplace, keep the evidence, and ask for help before a change becomes a violation."
)
p = doc.add_paragraph()
add_mixed(p, [
    "Related explainer: ",
    ("TFWP Employer Compliance", "https://commonwealthmigration.ca/work-and-study/tfwp-employer-compliance"),
    "."
])

doc.add_paragraph("Author bio", style="Heading 1")
doc.add_paragraph(
    "Anuj Dhawan is Content and Outreach Manager at Commonwealth Migration Group Inc., where he develops public-facing migration information for Canadian audiences."
)

doc.add_paragraph("Editorial note", style="Heading 1")
doc.add_paragraph(
    "This article is an informational, non-promotional submission. It is not legal advice, and employers should verify current program requirements and obtain advice on their specific facts."
)

doc.add_paragraph("AI use disclosure", style="Heading 1")
doc.add_paragraph(
    "AI assistance was used for brainstorming, outlining and language editing. The author reviewed the article, checked factual claims against the linked government sources and remains responsible for the final text."
)

doc.save(OUT)
print(OUT)
