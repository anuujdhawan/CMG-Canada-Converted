from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUTPUT = "outputs/refugee-research-online-article.docx"


def set_font(run, name="Aptos", size=11, bold=False, italic=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_font(run)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, size=14, bold=True)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
for style_name in ("Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.color.rgb = RGBColor(0, 0, 0)

# Remove Word's built-in title border so the submission has a clean editorial layout.
title_style = styles["Title"]
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(4)
title_run = title.add_run("Separating Official Requirements From Online Myth")
set_font(title_run, name="Aptos Display", size=22, bold=True)

dek = doc.add_paragraph()
dek.paragraph_format.space_after = Pt(10)
dek_run = dek.add_run("The information gap facing refugee claimants in Canada")
set_font(dek_run, size=12, italic=True, color="404040")

byline = doc.add_paragraph()
byline.paragraph_format.space_after = Pt(14)
byline_run = byline.add_run("By Anuj Dhawan")
set_font(byline_run, size=10, bold=True, color="404040")

add_body(doc, "For a person making a refugee claim in Canada, the hardest question is often not whether information exists. It is whether the information can be trusted, whether it applies to their situation, and whether it arrives in time to guide the next decision. Official instructions may sit beside community advice, old screenshots, translated summaries, social-media posts and messages from people who went through a different process. The result is an information environment that can be crowded and still leave a claimant unsure what to do.")

add_body(doc, "That problem is easy to describe as misinformation. It is better understood as an access problem. A right or procedure that cannot be understood, located and acted on is only partially available in practice. For refugee claimants, the consequences of confusion can include missed deadlines, incomplete forms, difficulty updating contact information, or failure to connect with legal and settlement support. These are not simply communication inconveniences. They can affect a person’s ability to participate in a high-stakes legal process.")

add_heading(doc, "A process with rules but no single doorway")
add_body(doc, "Canada’s refugee protection system is governed by formal rules and institutions. The Immigration and Refugee Board explains how a claim is made and how an eligible claim proceeds to the Refugee Protection Division. Immigration, Refugees and Citizenship Canada publishes information for claimants about forms, next steps and available support. Those public sources are essential, but they do not function as one universal doorway. A claimant may encounter them through a border official, a lawyer, a legal clinic, a settlement worker, a family member or a search engine, with each route carrying different context and urgency.")

add_body(doc, "The gap widens when a person’s circumstances change. A claimant may move, change a phone number, lose access to an email account, or rely on a family member to translate a notice. Someone who began with a lawyer may later need a clinic or community organization. Someone who arrived with a strong digital connection may lose it while moving between temporary homes. The formal process remains the same, but the claimant’s ability to reach and interpret the relevant information has changed.")

add_heading(doc, "The cost of contradiction")
add_body(doc, "People rarely rely on one source in isolation. They compare official pages with advice from friends, community groups and online forums because those sources answer different needs. An official page may state a rule accurately but assume that the reader already knows which institution is responsible. A community message may explain the issue in familiar language but repeat an outdated requirement. A legal professional may give precise advice that cannot be generalized to every claimant. The claimant is left to reconcile the sources, often without knowing which differences matter.")

add_body(doc, "This is why simply publishing more information is not enough. More pages can create more search results, more translations and more opportunities for contradiction. What matters is whether a claimant can identify the authoritative source, understand the date and scope of the information, and find a safe next step. Information quality includes provenance, plain language, accessibility and referral—not only factual accuracy in the narrow sense.")

add_heading(doc, "Design guidance around decisions")
add_body(doc, "A more useful information system would be organized around the decisions claimants actually face. Instead of asking a person to navigate a long list of institutions, a guide could begin with a question such as: Have you made a claim, are you preparing to make one, or have you received a notice about your claim? Each pathway could identify the official source, the relevant deadline or action, the kinds of support available, and the point at which a person should seek individualized legal advice.")

add_body(doc, "The same design principle applies to language access. Translation is important, but a translated page can still be difficult to use if its headings, dates, examples or referral links are unclear. Trusted intermediaries—legal clinics, settlement workers, libraries and community organizations—can help people test whether an explanation is understandable in the circumstances in which it will be used. Their role is not to replace official institutions. It is to make the route between an official requirement and a person’s next action more reliable.")

add_heading(doc, "A composite case study")
add_body(doc, "Consider a composite claimant in the Greater Toronto Area who makes an inland claim and receives instructions through several channels. The claimant finds the federal guidance online, receives a translated summary from a relative, speaks with a settlement worker about housing and health care, and later contacts a legal clinic after receiving a notice. None of these sources is necessarily wrong. The difficulty is that they do not share the same purpose, terminology or update schedule. The claimant must decide which source governs a deadline, which source explains a service, and which source is only general background.")

add_body(doc, "A practical intervention would connect these channels rather than treating them as separate information products. The official page could identify the decision it supports and link to a plain-language explanation from a trusted partner. The partner explanation could show its review date, distinguish general information from legal advice, and point back to the official source. A clinic or settlement organization could report recurring points of confusion without exposing confidential information. That feedback would help institutions improve the information at the point where people are actually getting stuck.")

add_heading(doc, "From information access to rights access")
add_body(doc, "The central question is not whether claimants are responsible for finding information. It is whether the surrounding system is designed for people who are under stress, may be working across languages, and may have limited time, connectivity or support. Treating information access as part of migration infrastructure changes the policy conversation. It asks institutions to measure not only whether a page exists, but whether people can find it, understand it, trust it and act on it without being pushed toward avoidable risk.")

add_body(doc, "For refugee claimants, better information will not solve every barrier. It cannot substitute for legal representation, fair decision-making or adequate settlement support. It can, however, reduce the distance between a formal right and the person trying to use it. The first step is to stop treating fragmented information as a minor communications problem. It is part of the practical conditions under which rights become real.")

add_heading(doc, "Selected public sources")
sources = [
    ("Immigration and Refugee Board of Canada, Applying for refugee protection", "https://irb-cisr.gc.ca/en/applying-refugee-protection/Pages/index.aspx"),
    ("Government of Canada, Information for refugee claimants", "https://www.canada.ca/en/immigration-refugees-citizenship/corporate/publications-manuals/information-refugee-claimants.html"),
]
for label, url in sources:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(3)
    add_hyperlink(p, label, url)

add_heading(doc, "Author bio")
add_body(doc, "Anuj Dhawan is Content and Outreach Manager at Commonwealth Migration Group, where he writes about newcomer information systems, settlement and public-facing migration communication. His work focuses on making complex migration processes easier to navigate through clear, sourced explanations.")

doc.core_properties.title = "Separating Official Requirements From Online Myth"
doc.core_properties.subject = "The information gap facing refugee claimants in Canada"
doc.core_properties.author = "Anuj Dhawan"
doc.save(OUTPUT)
print(OUTPUT)
