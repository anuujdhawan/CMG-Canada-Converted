from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


OUT = "outputs/policy-options-newcomers-settlement-op-ed.docx"


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_mixed(paragraph, pieces):
    for item in pieces:
        if isinstance(item, tuple):
            add_hyperlink(paragraph, item[0], item[1])
        else:
            paragraph.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08

article_title = styles.add_style("ArticleTitle", WD_STYLE_TYPE.PARAGRAPH)
article_title.base_style = styles["Normal"]
article_title.font.name = "Aptos Display"
article_title.font.size = Pt(22)
article_title.font.bold = True
article_title.font.color.rgb = RGBColor(0, 0, 0)
article_title.paragraph_format.space_after = Pt(5)
article_title.paragraph_format.keep_with_next = True

for name in ["Heading 1", "Heading 2"]:
    st = styles[name]
    st.font.name = "Aptos"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(5)

doc.add_paragraph("Newcomers need a clearer path through Canada’s settlement system", style="ArticleTitle")
byline = doc.add_paragraph()
byline.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = byline.add_run("By Anuj Dhawan")
run.bold = True
run.font.size = Pt(11)

doc.add_paragraph(
    "A family arriving in Brampton can find help with housing, language, work, health care and schools. The harder question is often where to begin. Information is available, but it is spread across federal directories, regional portals, municipal services, schools, libraries and community organizations. The result is a settlement system that can be generous in its offerings while still feeling difficult to navigate."
)

doc.add_paragraph(
    "Canada should measure the first stage of settlement not only by whether a newcomer was offered a service, but by whether that person reached the right service in time. Brampton and Peel Region could be a useful place to test a simple policy shift: organize information around the newcomer’s journey, then make every referral a supported handoff."
)

doc.add_paragraph("The information exists but the route is unclear", style="Heading 1")
doc.add_paragraph(
    "The federal settlement system already recognizes that newcomers need more than a list of programs. The Government of Canada describes settlement services as including needs assessment and referrals, information and orientation, language training, employment-related help and community connections. Its public service finder directs eligible newcomers to local providers and asks people to contact organizations before visiting."
    + " Those are sensible safeguards, but they also reveal the practical burden placed on the person seeking help: understand eligibility, identify the right category, find a provider, confirm that the provider offers the needed service and then arrange contact."
)

p = doc.add_paragraph()
add_mixed(p, [
    "The federal description of the Settlement Program is available in ",
    ("IRCC’s funding and service overview", "https://www.canada.ca/en/immigration-refugees-citizenship/corporate/partners-service-providers/funding.html"),
    ", while the ",
    ("newcomer service finder", "https://www.ircc.canada.ca/english/newcomers/services/index.asp"),
    " shows how many kinds of local support sit behind that national system."
])

doc.add_paragraph(
    "In Peel, the local picture is even more distributed. The Immigration Peel portal groups information about health care, schools, housing, banking, employment and community services, and lists newcomer information centres in Brampton. This network is valuable precisely because it is local. But a newcomer may still have to decide whether a question belongs with a settlement worker, an employment provider, a school welcome centre, a health service, a community organization or a government office."
)

p = doc.add_paragraph()
add_mixed(p, [
    "The ",
    ("Peel newcomer services directory", "https://www.immigrationpeel.ca/before-you-arrive/newcomer-services-in-peel-region/"),
    " illustrates the breadth of the local network. Its usefulness would be greater if every entry also made the next step, eligibility and handoff to another service unmistakably clear."
])

doc.add_paragraph("A changing system makes signposting more important", style="Heading 1")
doc.add_paragraph(
    "Eligibility rules add another layer of complexity. IRCC now explains that, starting April 1, 2026, economic-class permanent residents can access federally funded newcomer services for a defined period, while several other groups have different rules or no time limit. A person who arrives with an approval letter may reasonably assume that every settlement service is available in the same way. That assumption can be wrong, and a general web search rarely explains the difference in plain language."
)

p = doc.add_paragraph()
add_mixed(p, [
    "The current eligibility explanation is set out in ",
    ("IRCC’s guidance on newcomer-service eligibility", "https://www.canada.ca/en/immigration-refugees-citizenship/services/settle-canada/newcomer-services/eligibility.html"),
    ". The policy is understandable when read carefully; the challenge is making it usable at the moment a newcomer needs to act."
])

doc.add_paragraph("What a better first 90 days could look like", style="Heading 1")
doc.add_paragraph(
    "First, governments and service networks should publish a short, status-sensitive first-90-days map. It should begin with the newcomer’s situation—permanent resident, protected person, temporary worker, student or family member—not with the name of an agency. It should answer five questions in order: what can I do now, what documents will I need, which service is free, what if I am not eligible, and who can help me with the next step?"
)
doc.add_paragraph(
    "Second, referral quality should become a measurable outcome. A referral is not complete when a worker hands over a phone number. It is complete when the receiving service has accepted the referral, the newcomer knows what will happen next and there is a safe way to correct a failed handoff. Consent-based referral tools and simple follow-up messages could reduce the number of people who are repeatedly asked to retell the same story."
)
doc.add_paragraph(
    "Third, local directories should display freshness. Every service entry should show when its hours, eligibility and language information were last checked. A newcomer who receives an outdated address or discovers that a program has closed is not merely inconvenienced; trust in the wider settlement system is damaged."
)
doc.add_paragraph(
    "Finally, newcomers should help design the map. Libraries, schools, faith communities and settlement organizations already see where official language fails. Short user tests in Punjabi, Hindi, Urdu, Arabic, French and other languages commonly spoken in Brampton could reveal confusing terms, missing routes and assumptions that professionals no longer notice."
)

doc.add_paragraph("Brampton can demonstrate the change", style="Heading 1")
doc.add_paragraph(
    "Brampton does not need another isolated portal. It needs a shared agreement among the organizations that already serve newcomers: use plain language, identify eligibility, show the next step, keep information current and close the loop when a referral fails. The model could begin with a small set of high-pressure needs—housing, health care, school registration, employment and immigration status—and then be evaluated by newcomers themselves."
)
doc.add_paragraph(
    "The policy question is not whether Canada has enough information. It is whether a person can turn that information into a timely, safe decision. A clearer path through the first 90 days would make existing investments work better and give communities a practical way to see where the system still breaks. If Brampton can build that path locally, other Canadian communities could adapt it without starting from scratch."
)

doc.add_paragraph("Author bio", style="Heading 1")
doc.add_paragraph(
    "Anuj Dhawan is Content and Outreach Manager at Commonwealth Migration Group Inc., where he develops public-facing migration information for Canadian newcomer audiences."
)

doc.add_paragraph("AI use disclosure", style="Heading 1")
doc.add_paragraph(
    "AI assistance was used for brainstorming, outlining and language editing. The author reviewed the article, checked the factual claims against the linked sources and remains responsible for the final text."
)

doc.save(OUT)
print(OUT)
