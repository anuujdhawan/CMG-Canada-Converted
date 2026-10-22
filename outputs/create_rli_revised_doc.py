from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "outputs/rli-canada-case-study-revised.docx"


def set_run_font(run, name="Aptos", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_source(paragraph, label, url):
    paragraph.add_run(label + " ")
    add_hyperlink(paragraph, url, url)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08

for style_name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph("When the Right Exists on Paper", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle = doc.add_paragraph("A Canadian case study of the information gap after an inland refugee claim")
subtitle.paragraph_format.space_after = Pt(14)
set_run_font(subtitle.runs[0], size=12, italic=True, color=(80, 80, 80))

byline = doc.add_paragraph()
set_run_font(byline.add_run("By Anuj Dhawan"), bold=True)
byline.add_run(" | Content and Outreach Manager, Commonwealth Migration Group")

intro = doc.add_paragraph()
intro.add_run(
    "This article uses a composite, analytical case study based on publicly available Canadian guidance. "
    "It does not describe an individual client, rely on confidential information or present new interview data."
)

doc.add_paragraph("The access problem", style="Heading 1")
doc.add_paragraph(
    "A refugee claimant can have a legal right on paper and still struggle to use it in practice. "
    "The difficulty may not be a missing rule. It may be the distance between the rule and the moment when a person must act: "
    "submit a form, attend an interview, find counsel, arrange a medical examination, update an address or understand which document proves eligibility for a service."
)
doc.add_paragraph(
    "In Canada, the information required for that journey is distributed across several institutions. "
    "The Immigration, Refugees and Citizenship Canada (IRCC) process, the Canada Border Services Agency (CBSA), the Immigration and Refugee Board of Canada (IRB), provincial legal-aid systems and community organisations each provide part of the route. "
    "The information can be accurate in each location and still be difficult to use as a single pathway."
)
doc.add_paragraph(
    "This matters because a claimant is not only trying to understand refugee law. The claimant is also trying to work out what to do next, in which order, with which document and by which deadline. "
    "The question is therefore not simply whether official information exists. It is whether the system makes the next necessary decision visible to a person who may be new to Canada, unfamiliar with Canadian institutions, working in a second language or managing urgent housing and health needs at the same time."
)

doc.add_paragraph("A Canadian case study", style="Heading 1")
doc.add_paragraph(
    "Consider a claimant who makes an inland refugee claim after arriving in Canada and is settling in the Greater Toronto Area. "
    "The scenario is composite: it brings together the sequence and information demands described in public government and tribunal guidance rather than reconstructing any one person’s file. "
    "Its value is that it shows where a right can become difficult to access even when each institution has published relevant instructions."
)
doc.add_paragraph(
    "The first step is knowing which process applies. An inland claimant uses the IRCC Portal and must provide information and supporting documents. "
    "The claimant will also be contacted for biometrics and an interview. If the claim is eligible, it is referred to the IRB’s Refugee Protection Division. "
    "The IRB then describes the process through a sequence of separate steps: making the claim, sending the Basis of Claim form, preparing for the hearing, collecting and sending evidence, attending the hearing and receiving a decision."
)
doc.add_paragraph(
    "The information is not merely descriptive. The Basis of Claim form is mandatory, and the government’s claimant information explains that a port-of-entry claimant must submit it to the IRB within 15 days of referral, subject to the applicable instructions. "
    "The same guidance says that an inland claimant uploads the form through the IRCC Portal. A reader who does not understand the difference between the two routes may see two apparently similar instructions and still be unsure which deadline and submission channel applies."
)
doc.add_paragraph(
    "At the same time, the claimant needs help with issues that sit beside the legal determination. The claimant may need to understand whether an Acknowledgement of Claim or Refugee Protection Claimant Document can be used to access services, how Interim Federal Health Program coverage works, whether a medical examination is required, and how to keep IRCC, CBSA and the IRB informed after moving. "
    "The federal guidance connects these issues, but the claimant must still move between pages, institutions and forms to assemble a usable plan."
)

doc.add_paragraph("Where the handoffs become a barrier", style="Heading 1")
doc.add_paragraph(
    "The first handoff is from intake to the tribunal. A claimant may leave an intake interaction with several documents, but the documents do not all answer the same question. "
    "One confirms that a claim has been made; another indicates whether the claim has been referred; a notice later identifies the hearing; and the Basis of Claim form records the substance of the protection claim. "
    "A person who is reading under pressure may reasonably treat these documents as interchangeable even though they have different functions."
)
doc.add_paragraph(
    "The second handoff is from legal entitlement to representation. Canadian guidance explains that a claimant has the right to counsel at the claimant’s own expense and may apply for provincial or territorial legal aid if unable to pay. "
    "That is important information, but it does not by itself tell a claimant which legal-aid office serves the claimant’s location, what documents to bring, how urgently to seek help or what to do if a hearing-related deadline is approaching. "
    "A referral that says only ‘seek legal aid’ transfers the navigation burden to the person who is least likely to know the system."
)
doc.add_paragraph(
    "The third handoff is from status to everyday services. The Acknowledgement of Claim or a later claimant document may show eligibility for Interim Federal Health Program coverage and may help with access to services. "
    "A claimant may also need language support, housing information, a family doctor, school information for children or employment guidance. These services are delivered through different provincial, municipal and community channels. "
    "The legal right and the service pathway are connected in the claimant’s life, but they are not always connected in the way information is presented."
)
doc.add_paragraph(
    "The fourth handoff is the maintenance of contact information. The IRB warns that it must be able to reach claimants about hearings and important information, and its process page directs claimants to report changes. "
    "Federal guidance separately tells claimants to update IRCC and, after referral, the IRB. For someone who changes accommodation soon after arrival, the operational question is not simply ‘keep your address current’; it is ‘which bodies must I notify, through which channel, and what confirmation should I retain?’"
)

doc.add_paragraph("Why fragmentation affects rights", style="Heading 1")
doc.add_paragraph(
    "These are not minor usability concerns. A missed deadline can trigger additional proceedings or place a claim at risk. A missed hearing or an outdated address can prevent a claimant from receiving information that is essential to the process. "
    "Confusion about health coverage can delay care. Confusion about representation can delay the search for counsel. Each problem is made more likely when the person must build a coherent pathway from separate sources while also dealing with displacement, language barriers and financial insecurity."
)
doc.add_paragraph(
    "The point is not that every public page should reproduce every other institution’s guidance. Duplication can itself create inconsistency. The better approach is deliberate signposting: explain what the page answers, identify the next responsible institution, show which document or deadline matters, state the date of review and provide a direct route to current primary guidance."
)

doc.add_paragraph("A more usable pathway", style="Heading 1")
doc.add_paragraph(
    "First, public-facing information should be organised around moments of decision rather than institutional boundaries. A claimant who has submitted an inland claim needs a short, ordered pathway: what was submitted, what happens next, which document to watch for, what must be done by a deadline, and where to obtain legal or settlement support. "
    "The pathway can link back to the full IRCC and IRB pages without forcing the reader to discover the relationship alone."
)
doc.add_paragraph(
    "Second, referrals should be specific. ‘Contact legal aid’ is less useful than the name of the relevant provincial service, its eligibility page, the documents commonly requested and a warning that urgent deadlines should be raised immediately. "
    "The same principle applies to health coverage, interpreters and settlement services. A referral should identify the service’s role and its limits so that a general-information resource is not mistaken for legal advice."
)
doc.add_paragraph(
    "Third, information should be tested with the people who use it. Plain language is not achieved only by shortening sentences. It requires testing whether readers can identify the next action, distinguish a deadline from an estimate, understand which authority is responsible and find the correct language or accessibility support. "
    "Community organisations and legal clinics can help identify where an apparently complete explanation still leaves a practical gap."
)
doc.add_paragraph(
    "Fourth, institutions should treat document continuity as part of access. A claimant may need to show the same status document to several services. Guidance should explain the document’s purpose, what information it proves and what to do when a document is lost, expired or replaced. That continuity reduces repeated explanations and helps prevent a person from being sent from one service to another without a clear answer."
)

doc.add_paragraph("Conclusion", style="Heading 1")
doc.add_paragraph(
    "The Canadian case study shows how an information gap can emerge without a single institution providing incorrect guidance. The gap appears at the joins: between IRCC and the IRB, between a status document and a service, between a general referral and a time-sensitive legal need, and between a change of address and the several bodies that must be notified."
)
doc.add_paragraph(
    "Closing that gap is a practical part of refugee-rights protection. The goal is not to replace legal advice or simplify a process that has legitimate safeguards. It is to make the route through the process more intelligible, especially for people who must act before they have secure housing, stable income, a trusted adviser or confidence in the language of the system. A right is more meaningful when the person entitled to it can identify the next step and reach the institution responsible for it."
)

doc.add_paragraph("Sources and further reading", style="Heading 1")
sources = [
    ("Immigration and Refugee Board of Canada, The refugee protection claim process.", "https://irb-cisr.gc.ca/en/applying-refugee-protection/Pages/index.aspx"),
    ("Government of Canada, Information for Refugee Claimants.", "https://www.canada.ca/en/immigration-refugees-citizenship/corporate/publications-manuals/information-refugee-claimants.html"),
    ("Government of Canada, Guide 0174 for inland refugee claims submitted through the IRCC Portal.", "https://www.canada.ca/en/immigration-refugees-citizenship/services/application/application-forms-guides/guide-0174-inland-refugee-claims-portal.html"),
    ("Refugee Law Initiative, contribution guidelines.", "https://rli.blogs.sas.ac.uk/contribute/"),
    ("Commonwealth Migration Group, information on what happens after an in-Canada refugee claim.", "https://commonwealthmigration.ca/immigrate/in-canada-refugee-claim"),
]
for label, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    add_source(p, label, url)

doc.add_paragraph("Author bio", style="Heading 1")
doc.add_paragraph(
    "Anuj Dhawan is Content and Outreach Manager at Commonwealth Migration Group. His work focuses on making complex immigration and settlement information clearer, better sourced and easier for newcomers to navigate."
)

doc.add_paragraph("Submission note", style="Heading 1")
doc.add_paragraph(
    "This draft is original, evidence-led and written for a multidisciplinary policy audience. The case study is explicitly composite and analytical; it does not rely on interviews, participant observation or confidential client information. Commonwealth Migration Group is disclosed as the author’s professional affiliation. The practitioner resource in the sources is included as background and may be removed or edited to comply with the editorial policy of the Refugee Law Initiative."
)

doc.core_properties.title = "When the Right Exists on Paper"
doc.core_properties.subject = "Canadian case study of information access after an inland refugee claim"
doc.core_properties.author = "Anuj Dhawan"
doc.core_properties.keywords = "refugee law, Canada, information access, inland refugee claim, rights"
doc.save(OUT)
print(OUT)
