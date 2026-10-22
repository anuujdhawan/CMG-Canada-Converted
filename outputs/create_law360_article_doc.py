from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "outputs/law360-canada-article-review-copy.docx"


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(8)
styles["Normal"].paragraph_format.line_spacing = 1.08

for style_name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)

doc.add_paragraph("The information gap before access", style="Title")
subtitle = doc.add_paragraph("Turning Canadian immigration policy into usable guidance")
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle.runs[0].italic = True
subtitle.runs[0].font.size = Pt(12)
subtitle.runs[0].font.color.rgb = RGBColor(80, 80, 80)

meta = doc.add_table(rows=3, cols=2)
meta.style = "Table Grid"
meta.autofit = False
meta.columns[0].width = Inches(1.45)
meta.columns[1].width = Inches(5.9)
metadata = [
    ("Prepared for", "Law360 Canada editorial consideration"),
    ("Author", "Anuj Dhawan, Commonwealth Migration Group"),
    ("Format", "Review copy; original, non-promotional policy analysis"),
]
for row, (label, value) in zip(meta.rows, metadata):
    row.cells[0].text = label
    row.cells[1].text = value
    for cell in row.cells:
        set_cell_margins(cell)
    row.cells[0].paragraphs[0].runs[0].bold = True
    row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(45, 45, 45)

doc.add_paragraph()

sections = [
    (None, "Canadian immigration policy is often discussed as a question of eligibility: who qualifies, which stream is open, and what evidence must be submitted. For many newcomers, however, the practical problem begins after they have found the relevant program. They still need to understand what the rule means for work, family, housing, health care and the next administrative step. The gap between a policy being published and a person being able to use it is an access-to-justice issue."),
    (None, "The difficulty is not simply that immigration rules are complex. It is that the information needed to act is distributed across different institutions. A newcomer may read an Immigration, Refugees and Citizenship Canada page, consult a provincial service, speak with an employer, search a community forum and contact a legal professional, all within the same week. Each source may be accurate within its own scope, but the reader has to reconcile differences in language, timing and responsibility."),
    (None, "That reconciliation work creates avoidable risk. A person may confuse an eligibility requirement with a document checklist, treat a processing estimate as a deadline, or assume that a work-authorisation rule applies to a family member in the same way. The mistake may not come from ignoring the law. It may come from not knowing which institution has the authority to answer the question."),
    ("Better signposting", "Public guidance should make the next decision visible: whether the reader needs an eligibility explanation, an application instruction, an employment-rights resource, a settlement service or individual legal advice. A useful page should identify the responsible authority, the date it was last reviewed, the facts that may change the answer and the point at which a reader should stop relying on general information."),
    ("Information and advice", "Plain-language explainers are valuable when they translate technical requirements without implying that one generic answer resolves every case. They should distinguish official rules from practical suggestions, identify uncertainty and link readers back to the current government or tribunal source. A practitioner resource explaining the stages after an in-Canada refugee claim, for example, can help a reader prepare better questions while making clear that it is not a decision or a substitute for representation."),
    ("Institutional coordination", "Immigration policy does not operate in isolation from employment, education, health or family services. When those systems publish disconnected instructions, newcomers are left to build their own route through them. Cross-links, consistent terminology and referral agreements can reduce repetition and help people reach the right service earlier. This is particularly important for people who face language, disability, digital-access or financial barriers."),
    ("The role of professionals", "Employers and regulated professionals also have a role. An employer should not rely on a general impression about a worker’s immigration status when a specific document or condition matters. A service provider should not present marketing language as official policy. Both should direct people to authoritative sources and encourage qualified advice when the facts are case-specific. Clear boundaries protect newcomers and make professional communication more credible."),
    (None, "The measure of good immigration information is therefore not its volume or search visibility. It is whether a person can move from a question to a reliable next step without paying for information that should be public, repeating their history unnecessarily or relying on an outdated online claim. That requires content that is current, source-linked, transparent about limits and designed around the decisions people actually face."),
    (None, "Immigration policy is ultimately experienced through ordinary decisions: whether someone can accept a job, renew a document, bring a family member, access a service or understand what happens next. Closing the information gap will not remove every legal or administrative barrier. It can, however, make the system more intelligible and reduce the distance between a right on paper and a right that a newcomer can realistically use."),
]

for heading, text in sections:
    if heading:
        doc.add_paragraph(heading, style="Heading 1")
    doc.add_paragraph(text)

doc.add_paragraph("Background reference", style="Heading 1")
doc.add_paragraph(
    "Commonwealth Migration Group, What Happens After a Refugee Claim Is Made in Canada: "
    "https://commonwealthmigration.ca/immigrate/in-canada-refugee-claim"
)

doc.add_paragraph("Author note", style="Heading 1")
doc.add_paragraph(
    "Anuj Dhawan is Content and Outreach Manager at Commonwealth Migration Group, where he works on public-facing immigration information and migration-policy content. His work focuses on making complex immigration and settlement information clearer, better sourced and easier for newcomers to navigate."
)

doc.add_paragraph("Editorial note", style="Heading 1")
doc.add_paragraph(
    "This review copy is original, evidence-led and non-promotional. Commonwealth Migration Group is disclosed as the author’s professional affiliation. Any external link may be removed or edited to comply with Law360 Canada’s normal editorial policy."
)

doc.core_properties.title = "The information gap before access"
doc.core_properties.subject = "Law360 Canada article review copy"
doc.core_properties.author = "Anuj Dhawan"
doc.save(OUT)
print(OUT)
