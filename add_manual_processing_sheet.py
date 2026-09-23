from openpyxl import load_workbook
from pathlib import Path

path = Path('outputs/backlink-created-details.xlsx')
wb = load_workbook(path)
name = 'Manual Processing'
if name in wb.sheetnames:
    ws = wb[name]
else:
    ws = wb.create_sheet(name)
    ws.append(['Website URL', 'Reason for manual processing', 'Evidence URL', 'Recorded date'])

row = [
    'https://todayincanada.ca/contact/',
    'reCAPTCHA checkbox required; human completion needed',
    'https://todayincanada.ca/contact/',
    '2026-09-23',
]
rows = [row, [
    'https://thej.ca/submit-article/',
    'reCAPTCHA protection present; human completion needed',
    'https://thej.ca/submit-article/',
    '2026-09-23',
]]
existing = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
for item in rows:
    if item not in existing:
        ws.append(item)
for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    for col in sheet.columns:
        width = max(len(str(c.value or '')) for c in col) + 2
        sheet.column_dimensions[col[0].column_letter].width = min(max(width, 18), 80)
wb.save(path)
print(path)
