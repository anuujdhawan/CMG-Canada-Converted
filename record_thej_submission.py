from openpyxl import load_workbook
from pathlib import Path

details = Path('outputs/backlink-created-details.xlsx')
wb = load_workbook(details)
ws = wb['Backlink Details']
row = ['https://thej.ca/submit-article/', 'https://commonwealthmigration.ca/immigrate/express-entry']
if row not in [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]:
    ws.append(row)
wb.save(details)

progress = Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb2 = load_workbook(progress)
log = wb2['Submission Log']
logrow = ['THEJ.CA', 'https://thej.ca/submit-article/', 'Direct article submission form', 'Submitted/pending', 'https://thej.ca/thank-you-for-getting-in-touch/', 'Form returned MESSAGE SENT / thank-you page. Article contains a contextual CMG Express Entry URL. No public article permalink or backlink publication has been verified.', 'Monitor for publication; do not count as published until publicly verified', '2026-09-23']
if tuple(logrow) not in {tuple(r) for r in log.iter_rows(min_row=2, values_only=True)}:
    log.append(logrow)
wb2.save(progress)
print(details)
print(progress)
