from openpyxl import load_workbook
from pathlib import Path
from datetime import date
p=Path('outputs/backlink-outreach-progress-2026-09-23.xlsx'); wb=load_workbook(p)
ws=wb['Prospects']
for r in range(2,ws.max_row+1):
    if ws.cell(r,1).value=='BetterPlace Immigration':
        ws.cell(r,10).value='Current page invites expert/feature pieces and provides info@betterplaceimmigration.com. No email sent: native Mail is configured only for an unrelated iCloud account, not the authorized seolaunchers@gmail.com mailbox.'
        ws.cell(r,12).value='Research verified 2026-09-23; blocked pending authorized mailbox'
wb['Submission Log'].append(['BetterPlace Immigration','https://www.betterplaceimmigration.com/write-for-us/','Email-only; editorial review','Blocked; not submitted','https://www.betterplaceimmigration.com/write-for-us/','Viable editorial route, but no authorized seolaunchers@gmail.com mailbox is configured. Native Mail is connected to an unrelated iCloud account; no message sent.','Configure/use the authorized mailbox before transmitting the pitch.',date(2026,9,23)])
wb.save(p); print(p.resolve())
