from openpyxl import load_workbook
from pathlib import Path
p=Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb=load_workbook(p)
ws=wb['Prospects']
for r in range(2, ws.max_row+1):
    if ws.cell(r,1).value=='CanConnect':
        ws.cell(r,8).value='Attempted; outcome unconfirmed'
        ws.cell(r,10).value='Contribution interest form accepted the content but remained stuck on “Submitting…” with no thank-you/receipt or navigation. Do not treat as submitted/pending or published; no backlink exists.'
        ws.cell(r,12).value='Attempted 2026-09-23; no confirmation returned'
log=wb['Submission Log']
log.append(['CanConnect','https://canconnect.ca/get-involved','Direct interest form; human review','Attempted; outcome unconfirmed','https://canconnect.ca/get-involved','Form remained on disabled “Submitting…” state after content was entered; no confirmation or receipt visible. No backlink exists.','Do not retry automatically; use the listed info@canconnect.ca email only through an authorized mailbox if follow-up is approved.',__import__('datetime').date(2026,9,23)])
wb.save(p)
print(p.resolve())
