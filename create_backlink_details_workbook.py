from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pathlib import Path

p=Path('/Users/themacintosh/Documents/My Files/Projects/CMG-Tailwind-components(V2)/outputs/backlink-created-details.xlsx')
if p.exists():
    wb=load_workbook(p)
    ws=wb.active
else:
    wb=Workbook()
    ws=wb.active
    ws.title='Backlink Details'
    ws.append(['Website URL where article was submitted','Company website backlink URL included in article'])
    for c in ws[1]: c.font=Font(bold=True)
    ws.freeze_panes='A2'
    ws.column_dimensions['A'].width=65
    ws.column_dimensions['B'].width=65
wb.save(p)
print(p, ws.max_row, ws.max_column)
