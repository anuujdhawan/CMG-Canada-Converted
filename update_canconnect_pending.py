from openpyxl import load_workbook
from pathlib import Path
from datetime import date
p=Path('outputs/backlink-outreach-progress-2026-09-23.xlsx'); wb=load_workbook(p)
ws=wb['Prospects']
for r in range(2,ws.max_row+1):
    if ws.cell(r,1).value=='CanConnect':
        ws.cell(r,8).value='Submitted/pending'
        ws.cell(r,10).value='CanConnect sent an automated acknowledgement from noreply@canconnect.ca: the message will be reviewed and a team member will follow up. No article is live and no backlink is verified.'
        ws.cell(r,12).value='Submission receipt verified in Gmail 2026-09-23; pending editorial review'
log=wb['Submission Log']
log.append(['CanConnect','https://canconnect.ca/get-involved','Direct interest form; human review','Submitted/pending','Gmail acknowledgement from noreply@canconnect.ca; no public article URL','Automated acknowledgement received: message will be reviewed and team will follow up. No backlink exists yet.','Wait for editorial follow-up; verify only if a public article is published.',date(2026,9,23)])
wb.save(p); print(p.resolve())
