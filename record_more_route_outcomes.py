from openpyxl import load_workbook
from pathlib import Path

path = Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb = load_workbook(path)
ws = wb['Submission Log']
rows = [
    ['Cyboard School', 'https://cyboardschool.com/write-for-us/', 'Not suitable: reciprocal-link requirement', 'Not submitted', 'https://cyboardschool.com/write-for-us/', 'The route requires a live backlink to Cyboard from the contributor website before publication and sets DA/PA/spam-score thresholds; no reciprocal arrangement was authorized or verified.', 'Exclude from current campaign', '2026-09-23'],
    ['Over Here Toronto', 'https://www.overheretoronto.com/submit-story/', 'Not suitable: editorial rule conflict', 'Not submitted', 'https://www.overheretoronto.com/submit-story/', 'Direct form is accessible, but the site requires original human-written content and prohibits AI content; the current campaign cannot truthfully satisfy that requirement.', 'Exclude from current campaign', '2026-09-23'],
    ['Fireside Publishing House', 'https://firesidepublishinghouse.ca/contribute.html', 'Not suitable: local connection required', 'Not submitted', 'https://firesidepublishinghouse.ca/contribute.html', 'Direct pitch form is available, but the publication requires a genuine connection to Kawartha Lakes and lived/local experience; a generic CMG article would not meet the stated fit.', 'Exclude from current campaign', '2026-09-23'],
]
existing = {tuple(r) for r in ws.iter_rows(min_row=2, values_only=True)}
for row in rows:
    if tuple(row) not in existing:
        ws.append(row)
wb.save(path)
print(path)
