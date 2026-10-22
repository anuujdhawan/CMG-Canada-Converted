from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "rli-canada-case-study-rebuilt.docx"


def set_run_font(run, name="Aptos", size=11, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A5FB4")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rpr.append(sz)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_mixed_paragraph(doc, pieces, style=None, space_after=7):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    for piece in pieces:
        if isinstance(piece, tuple) and piece[0] == "link":
            add_hyperlink(p, piece[1], piece[2])
        else:
            r = p.add_run(piece)
            set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 12, bold=True)
    return p


def add_bullet(doc, pieces):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    for piece in pieces:
        if isinstance(piece, tuple) and piece[0] == "link":
            add_hyperlink(p, piece[1], piece[2])
        else:
            r = p.add_run(piece)
            set_run_font(r)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.95)
section.right_margin = Inches(0.95)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

for style_name, size in [("Title", 19), ("Heading 1", 14), ("Heading 2", 12)]:
    st = styles[style_name]
    st.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    ppr = st._element.get_or_add_pPr()
    for child in list(ppr):
        if child.tag == qn("w:pBdr"):
            ppr.remove(child)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.keep_with_next = True
r = title.add_run("When Rights Are Hard to Reach")
set_run_font(r, name="Aptos Display", size=19, bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
subtitle.paragraph_format.keep_with_next = True
r = subtitle.add_run("A Canadian case study of the information gap after an inland refugee claim")
set_run_font(r, size=12, italic=True)

byline = doc.add_paragraph()
byline.paragraph_format.space_after = Pt(15)
r = byline.add_run("By Anuj Dhawan")
set_run_font(r, size=10.5, bold=True)

add_mixed_paragraph(doc, [
    "A person who makes a refugee claim in Canada does not encounter one system. They encounter a chain of systems: immigration intake, the Immigration and Refugee Board of Canada (IRB), legal aid, health coverage, employment services, schools, settlement agencies and, often, municipal or community support. Each institution may provide accurate information. The difficulty is that the information is usually delivered at a different time, through a different portal or document, with a different vocabulary and a different definition of what the person needs to do next."
])

add_mixed_paragraph(doc, [
    "This is an access problem rather than simply an information problem. A right can exist in law while remaining difficult to use if the claimant cannot identify the next step, cannot prove eligibility at the moment of service, or does not know which institution must be told about a change. The gap is particularly visible after an inland claim, when a claimant may be waiting for a hearing while also trying to secure housing, health care, work, legal representation and a stable address."
])

add_mixed_paragraph(doc, [
    "The following is a composite analytical case study based on public Canadian guidance. It does not describe a real client, interview or confidential file. Its purpose is to show how ordinary administrative handoffs can turn formal rights into uncertain, delayed or uneven access."
])

add_heading(doc, "A claimant arrives in Brampton", 1)
add_mixed_paragraph(doc, [
    "Consider the case of Samira, a composite claimant who makes an inland refugee claim after arriving in the Greater Toronto Area. She is living temporarily with relatives in Brampton. Her claim is eligible and is referred to the Refugee Protection Division (RPD). She receives several documents, is told to complete the next stages of the process, and is advised to obtain medical attention, seek legal advice and apply for work if she is eligible."
])

add_mixed_paragraph(doc, [
    "Nothing in that sequence is unusual. The problem is that the instructions do not arrive as one coherent pathway. The federal information for refugee claimants explains what documents mean, how medical examinations and work permits fit into the process, and how contact information must be updated. The IRB separately explains the RPD process, the Basis of Claim (BOC) form, evidence and hearing preparation. Ontario directs newcomers toward settlement agencies, health services and 211. Legal Aid Ontario describes its own eligibility rules and certificate process. Each guide is useful on its own. Together, they require the claimant to build the map."
])

add_heading(doc, "The first handoff is from a claim to a case file", 1)
add_mixed_paragraph(doc, [
    "The first risk appears at the point where a claim moves from immigration intake to the refugee determination process. The claimant may understand that the claim has been accepted for processing without understanding what the referral to the IRB changes. The RPD becomes the institution that must be able to contact the claimant, issue hearing information and receive documents. The claimant is responsible for keeping the RPD informed of changes in contact information; the IRB warns that an outdated address can lead to missed notices and, in turn, an abandonment process."
])

add_mixed_paragraph(doc, [
    "For Samira, the instruction to update contact information is not merely clerical. Her temporary accommodation ends before she has found permanent housing. She tells a settlement worker about the move, updates her address through an IRCC channel, and assumes the file is now updated everywhere. But the federal guidance says that, once a claim has been referred to the IRB, the claimant must also advise the IRB of the change. The two updates are related but not identical. A system that appears unified from the claimant's perspective is actually divided between institutions."
])

add_mixed_paragraph(doc, [
    "This is a classic information handoff. The relevant question is not whether an address-update form exists. It is whether the claimant knows which organisation holds which part of the record, whether the claimant can submit the update in time, and whether the update is confirmed in a way that can be shown later. A missed hearing notice is therefore not always the product of inattention. It can be the foreseeable result of fragmented responsibility."
])

add_heading(doc, "A work permit is not the same as access to work", 1)
add_mixed_paragraph(doc, [
    "The same fragmentation appears in employment. The Government of Canada explains that an eligible refugee claimant may request an open work permit and that a work permit is required to work legally. It also explains that the claimant needs a Social Insurance Number (SIN) once the permit is issued. For a person trying to pay rent and buy food, however, the practical pathway is longer: claim eligibility, medical examination, biometrics where required, receipt of the permit, SIN application, employer verification and, sometimes, a new address for the permit to be mailed."
])

add_mixed_paragraph(doc, [
    "Samira receives informal job leads before the work permit arrives. One employer asks for a SIN; another wants proof that she can work. She does not know whether the document she has is evidence of a pending application or evidence of current work authorization. The distinction matters. An open work permit can expand her options, but only once it is issued and valid. The legal right to work is therefore mediated by a document chain that the claimant must understand and keep intact."
])

add_mixed_paragraph(doc, [
    "A short, plain-language checklist could prevent much of this uncertainty: what document proves the claim was referred; what document proves work authorization; when a SIN can be requested; which address should be used for delivery; and what to do if a permit does not arrive. The existing guidance contains these pieces, but they are distributed across separate pages and institutions. The missing intervention is not another general explanation of the refugee system. It is a reliable handoff at the moment when a claimant moves from legal status to everyday administration."
])

add_heading(doc, "Health coverage depends on presenting the right proof", 1)
add_mixed_paragraph(doc, [
    "Health care provides another example of the difference between entitlement and usability. The Government of Canada states that the Interim Federal Health Program (IFHP) covers certain health-care services and products for eligible refugee claimants. Ontario's settlement information also directs refugee claimants to the IFHP, community health centres, settlement agencies and the Refugee HealthLine. These are important routes, but they are not interchangeable. A clinic, pharmacy or provider may need to see a current refugee protection identity document or other proof of coverage, while a claimant may know only that health care is supposed to be covered."
])

add_mixed_paragraph(doc, [
    "When Samira visits a clinic, the receptionist asks for a provincial health card. Samira does not have one. She presents the document from her claim and is told that the clinic needs to verify the coverage. She leaves without care, even though the relevant federal and provincial information describes a route to treatment. The failure is not necessarily a denial of the legal entitlement. It is a failure of translation between a federal status document and a front-line service encounter."
])

add_mixed_paragraph(doc, [
    "This is also where language and digital access become legal issues in practice. A claimant who cannot interpret an acronym, access an online portal or explain the difference between an immigration medical examination and treatment may delay care. A settlement agency can help bridge the gap, but the referral must arrive before the crisis. The most useful information is often not a broad statement that services exist; it is a sentence explaining what to carry, what to ask for and what to do when the first provider is unfamiliar with the program."
])

add_heading(doc, "Legal representation is a time-sensitive referral", 1)
add_mixed_paragraph(doc, [
    "The legal pathway has a similar structure. The IRB explains that claimants may be represented and that they are responsible for submitting material relevant to their claim. Legal Aid Ontario offers assistance in refugee and immigration matters, including help with the BOC and representation before the RPD, subject to its eligibility and certificate rules. The existence of a legal-aid program does not mean that representation is automatically in place. The claimant must find the program, understand the financial and procedural requirements, locate a lawyer or clinic and provide the information needed for an application."
])

add_mixed_paragraph(doc, [
    "For Samira, the referral comes in two directions at once. The RPD process requires her to prepare a coherent account and evidence. Legal aid asks for financial and personal information to determine eligibility. A settlement worker can help her find a service, but cannot replace legal advice. If the referral is late, or if the claimant assumes that an initial consultation is the same as representation, the consequences may appear later as missing evidence, confusion about deadlines or an inability to explain a change in circumstances."
])

add_mixed_paragraph(doc, [
    "The practical lesson is that legal information should be delivered as a sequence of decisions rather than as a directory of services. At each stage, a claimant should be able to answer four questions: what has happened to my claim; what must I do next; what document proves my position; and who must be notified if something changes? The answers are not all held by one institution, but the claimant experiences them as one case."
])

add_heading(doc, "What the Brampton case study reveals", 1)
add_mixed_paragraph(doc, [
    "Samira's experience illustrates four recurring features of the information gap. First, the system is institutionally fragmented. IRCC, the IRB, Ontario programs, Legal Aid Ontario and community organisations each have legitimate responsibilities, but the claimant must coordinate the boundaries between them."
])
add_bullet(doc, ["Second, the system is temporally fragmented. A work permit, medical exam, BOC deadline, hearing notice and legal-aid application do not arrive at the same moment. Information that is technically available may not be available when the claimant needs to act."])
add_bullet(doc, ["Third, the system is document dependent. A claimant may be eligible for a service but unable to prove eligibility to a front-line provider, especially when a status document has changed or a record has not yet been updated."])
add_bullet(doc, ["Fourth, the system is digitally uneven. Online forms and portals can speed up a process for a claimant with stable internet, a private device, strong English or French and confidence with official terminology. They can create a further barrier for someone relying on a shared phone, public Wi-Fi, translated assistance or an informal address."])

add_mixed_paragraph(doc, [
    "These features matter because each small uncertainty can compound. A move can lead to a missed notice. A missed notice can delay a hearing. A delayed hearing can affect work, housing and legal planning. A missing document can turn a covered health service into an out-of-pocket expense or an abandoned appointment. The system does not need to contain a single unlawful rule for a claimant to experience a serious practical loss."
])

add_heading(doc, "From signposting to continuity", 1)
add_mixed_paragraph(doc, [
    "The policy response should focus on continuity between institutions. A first step would be a claimant-facing timeline that follows the person's experience rather than the government's organisational chart. It could show the claim, referral, BOC, medical examination, work permit, SIN, legal aid, health coverage, address updates and hearing preparation in the order in which a claimant is likely to encounter them. Each item should identify the responsible institution, the document to keep and the action required if circumstances change."
])

add_mixed_paragraph(doc, [
    "A second step would be a standard handoff record. With the claimant's consent, a settlement agency or other authorised support organisation could give the claimant a dated checklist recording which institution has been contacted, what confirmation was received and what remains outstanding. This would not create a new immigration file. It would create a portable explanation of the claimant's administrative position, reducing the need to retell the same story at every service point."
])

add_mixed_paragraph(doc, [
    "A third step would be better front-line translation. Clinics, employers, schools, libraries and settlement workers do not need to become immigration specialists. They do need short guidance on the documents commonly presented by refugee claimants, the limits of each document and the correct referral route. The goal is not to outsource legal responsibility. It is to stop a claimant from being sent away simply because the first worker does not recognise the proof of status."
])

add_mixed_paragraph(doc, [
    "Finally, feedback should travel in both directions. If a settlement agency repeatedly sees claimants miss a particular IRB notice, or if clinics repeatedly encounter confusion about IFHP documents, the pattern should be visible to the institutions that issue the information. A service system that only sends instructions outward but does not receive evidence about how those instructions work in practice will continue to mistake publication for access."
])

add_heading(doc, "Conclusion", 1)
add_mixed_paragraph(doc, [
    "The Brampton case study is not an argument that Canada lacks refugee-protection rules or public information. It is an argument that rights become usable only when the institutions around a claimant connect their information at the right time and in a form that can be acted on. For an inland refugee claimant, the journey from arrival to a hearing is also a journey through documents, portals, referrals and proofs."
])

add_mixed_paragraph(doc, [
    "Closing the information gap therefore requires more than clearer webpages. It requires continuity: a shared understanding of the next step, a reliable way to update contact information, documents that travel with the claimant, and front-line services prepared to recognise the evidence of status. The legal question is not only whether a right exists on paper. It is whether the claimant can reach it before the deadline, the missed appointment or the next change of address."
])

add_heading(doc, "Sources and further reading", 1)
refs = [
    ("Government of Canada, Information for Refugee Claimants", "https://www.canada.ca/en/immigration-refugees-citizenship/corporate/publications-manuals/information-refugee-claimants.html"),
    ("Government of Canada, Claiming asylum from within Canada while you wait", "https://www.canada.ca/en/immigration-refugees-citizenship/services/asylum/in-canada/while-you-wait.html"),
    ("Immigration and Refugee Board of Canada, The refugee protection claim process", "https://irb-cisr.gc.ca/en/applying-refugee-protection/Pages/index.aspx"),
    ("Immigration and Refugee Board of Canada, Claimant's guide", "https://irb-cisr.gc.ca/en/refugee-claims/Pages/ClaDemGuide.aspx"),
    ("Ontario, Getting settled in Ontario", "https://www.ontario.ca/page/getting-settled-ontario"),
    ("Legal Aid Ontario, Refugee law", "https://www.legalaid.on.ca/faq/refugee-law/"),
    ("Legal Aid Ontario, Refugee Law Office", "https://www.legalaid.on.ca/more/corporate/contact-legal-aid-ontario/refugee-law-office/"),
]
for label, url in refs:
    add_bullet(doc, [("link", label, url)])

add_heading(doc, "Author bio", 1)
add_mixed_paragraph(doc, [
    "Anuj Dhawan is Content and Outreach Manager at Commonwealth Migration Group Inc. His work focuses on explaining Canadian immigration and settlement processes in clear, practical language. He is interested in the relationship between public information, administrative systems and the ability of newcomers and refugee claimants to use the rights and services available to them."
])

add_heading(doc, "Submission note", 1)
add_mixed_paragraph(doc, [
    "I am submitting this article because the practical accessibility of refugee-protection systems is often discussed as a communications issue when it is also a question of rights, timing and institutional design. The article uses a composite case study and public sources only; it does not rely on confidential client information, interviews, participant observation or new research data."
])

doc.save(OUT)
print(OUT)
