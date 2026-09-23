from openpyxl import load_workbook
p='/Users/themacintosh/Documents/My Files/Projects/CMG-Tailwind-components(V2)/outputs/backlink-outreach-progress-2026-09-23.xlsx'
wb=load_workbook(p)
ws=wb['Backlink Submission URLs']
row=['https://ceris.ca/write-for-us/','https://commonwealthmigration.ca/immigrate/express-entry']
if row not in [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]: ws.append(row)
log=wb['Submission Log']
log.append(['CERIS News Blog Resources','https://ceris.ca/write-for-us/','Direct article/contact form; no login','Attempted; outcome unconfirmed','https://ceris.ca/write-for-us/','Original article with contextual CMG Express Entry URL entered into the direct form. The form remained on the same page and showed no receipt or confirmation.','Do not claim publication or a live backlink; monitor for response/public URL.',__import__('datetime').datetime(2026,9,23)])
wb.save(p)
print('recorded',row,'rows',ws.max_row,'log',log.max_row)
