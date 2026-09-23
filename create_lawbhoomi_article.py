from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('outputs/lawbhoomi-immigration-article.docx')

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    props.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)
    run.append(props)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)

def add_body(doc, before, link_text=None, link_url=None, after=''):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    p.add_run(before)
    if link_text and link_url:
        add_hyperlink(p, link_text, link_url)
    p.add_run(after)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)

normal = doc.styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)

for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

doc.add_heading('How Newcomers Can Build a Reliable Canadian Immigration Information Plan', 0)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run('By Anooj Dhawan  |  Immigration and relocation information')
r.italic = True
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(85, 85, 85)

add_body(doc, 'Immigration decisions often begin with a simple online search, but the first answer a person finds is not always current, complete, or meant for their situation. A reliable information plan helps applicants move from general interest to careful preparation. It also reduces the risk of paying for advice before understanding what question needs to be answered.')

def h(text):
    doc.add_heading(text, 1)

h('Start with the decision you are actually making')
add_body(doc, 'The word immigration covers many different decisions. A person may be comparing permanent residence programs, preparing a work permit application, looking at study options, planning a family visit, or checking whether a credential can be used in Canada. Each question has different rules and evidence. Write the decision in one sentence before searching. For example: “I want to know which documents I need to assess my education for an economic immigration application.” A precise question produces better research than a broad search for “easy Canada visa.”')

h('Use official sources as the control document')
add_body(doc, 'Government pages should be the control document for eligibility, fees, forms, deadlines, document requirements, and application steps. Independent articles can explain a complicated process, but they should not replace the source that controls the outcome. Save the official URL, the date you read it, and the specific statement that answers your question. If two explanations conflict, return to the official source and check whether one page is older or applies to a different program.')

h('Separate a general overview from personal advice')
add_body(doc, 'A guide can explain how a program works without deciding whether a particular person qualifies. Eligibility may depend on details such as work history, duties, language results, education, family composition, location, status, and previous applications. A checklist copied from another applicant may therefore be incomplete. Treat general articles as a map of the questions you need to investigate, not as a personal assessment or a promise of approval.')

h('Build a source log before collecting documents')
add_body(doc, 'A simple source log can prevent confusion later. Create columns for the topic, source URL, date checked, key requirement, documents mentioned, and follow-up question. Add an entry whenever you read an official page, a provincial instruction, or a reputable settlement resource. Record the exact program name and version of any form. This practice is especially useful when family members are researching different pathways at the same time.')

h('Understand the role of Express Entry information')
add_body(doc, 'People exploring economic immigration often encounter Express Entry early in their research. A plain-language overview can help explain the broad structure, but applicants still need to compare their own education, language, work experience, and settlement plans with current government instructions. A useful starting point is ', 'Commonwealth Migration Group’s Express Entry overview', 'https://commonwealthmigration.ca/immigrate/express-entry', '. It is an orientation resource, not a substitute for official requirements or individualized legal advice. Readers should verify every current rule and form through Government of Canada sources before acting.')

h('Check who is giving the advice')
add_body(doc, 'Before relying on a website, identify its author, publication date, update date, and purpose. A government page, a university resource, a settlement organization, a law-firm explanation, and a commercial service page may each have a different role. Commercial interests do not automatically make information wrong, but they should be disclosed and considered. Be cautious with anonymous claims, copied checklists, guaranteed outcomes, pressure to pay immediately, or instructions that conflict with official sources.')

h('Ask for evidence instead of certainty')
add_body(doc, 'Good immigration information explains what is known, what may change, and where the reader should verify the detail. Processing times, fees, invitation thresholds, program openings, and document rules can change. An article that gives a date or number without a source may become misleading quickly. Ask: What is the source? When was it updated? Does it apply to my program and location? What would make this information no longer accurate?')

h('Protect personal information while researching')
add_body(doc, 'Research does not require sending a passport, bank statement, or full family history to every website that offers a checklist. Share only the information needed for the immediate question, and use secure channels for sensitive documents. Never provide account passwords or one-time verification codes to someone who claims they can speed up an application. Keep a private record of payments, contracts, correspondence, and documents that were sent.')

h('Turn research into a decision checklist')
add_body(doc, 'At the end of the research stage, write a short checklist: the pathway being considered, the official eligibility page, the required language or education evidence, the work-history questions that remain, the estimated costs to verify, and the next action. If an important question is unanswered, pause before submitting anything. A short pause for verification is usually safer than rebuilding an application after relying on an outdated or incomplete explanation.')

h('A careful plan is more useful than a confident shortcut')
add_body(doc, 'Reliable immigration research is not about finding the most certain-sounding answer. It is about matching the right question to the right source, recording what was checked, and recognizing when personal circumstances require qualified advice. Newcomers and prospective applicants can make better decisions when they keep official guidance at the centre, use independent explanations for context, and treat commercial information transparently. That process takes time, but it creates a clearer record and a stronger foundation for whatever decision comes next.')

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(0)
r = p.add_run('Author note: ')
r.bold = True
p.add_run('This article provides general information and does not constitute legal advice or a promise of an immigration outcome. Readers should confirm current requirements through official government sources.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
