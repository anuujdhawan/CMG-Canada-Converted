from openpyxl import load_workbook
from pathlib import Path
p=Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb=load_workbook(p)
ws=wb['Prospects']
for r in range(2, ws.max_row+1):
    if ws.cell(r,1).value=='Canooq':
        ws.cell(r,8).value='Attempted; outcome unconfirmed'
        ws.cell(r,10).value='Contribution pitch entered and Send message activated. Site navigated to a query-string URL containing the form fields, but no thank-you/receipt or server confirmation was visible. Do not treat as submitted/pending or published; no backlink exists.'
        ws.cell(r,12).value='Attempted 2026-09-23; evidence URL recorded in Submission Log'
log=wb['Submission Log']
for r in range(2, log.max_row+1):
    if log.cell(r,1).value=='Canooq':
        log.cell(r,4).value='Attempted; outcome unconfirmed'
        log.cell(r,5).value='https://www.canooq.ca/contact?website=&name=Commonwealth+Migration+Group&email=seolaunchers%40gmail.com&topic=Other&message=Subject%3A+A+practical+pre-arrival+checklist+for+newcomers+moving+to+Canada'
        log.cell(r,6).value='Send action navigated to a query-string URL with the entered fields; no thank-you/receipt or server confirmation visible. No backlink exists.'
        log.cell(r,7).value='Do not retry automatically; follow up by email only if an authorized mailbox is available.'
wb.save(p)
print(p.resolve())
