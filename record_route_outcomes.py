from openpyxl import load_workbook
from pathlib import Path

path = Path('outputs/backlink-outreach-progress-2026-09-23.xlsx')
wb = load_workbook(path)
ws = wb['Submission Log']
rows = [
    ['Orbus Canada', 'https://orbus.ca/write-for-us/', 'Blocked before form', 'Not submitted', 'https://orbus.ca/write-for-us/', 'Guest-post rules are relevant, but the contact form is embedded in an inaccessible external frame in the browser session; no safe field-level submission was possible.', 'Do not claim submission; revisit with an accessible form', '2026-09-23'],
    ['Canada Careers Conversation', 'https://canadacareersconversation.com/submit-content', 'Blocked before form', 'Not submitted', 'https://canadacareersconversation.com/submit-content', 'Submission page loaded blank in the browser session after waiting; no accessible fields or confirmation were available.', 'Do not claim submission; revisit if the form becomes accessible', '2026-09-23'],
]
existing = {tuple(r) for r in ws.iter_rows(min_row=2, values_only=True)}
for row in rows:
    if tuple(row) not in existing:
        ws.append(row)
wb.save(path)
print(path)
