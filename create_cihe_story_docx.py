from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

out = Path('outputs/cihe-mary-ann-shadd-cary-story.docx')
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(8)
styles['Normal'].paragraph_format.line_spacing = 1.08
styles['Title'].font.name = 'Aptos Display'
styles['Title'].font.size = Pt(20)
styles['Title'].font.bold = True
styles['Title'].font.color.rgb = None

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run('Mary Ann Shadd Cary and the Work of Building Canada')

intro = doc.add_paragraph()
intro.add_run('Submission context. ').bold = True
intro.add_run('This short historical reflection is prepared for the Canadian Institute for Historical Education’s Telling Canada’s Stories campaign. It focuses on an immigrant and newcomer whose education, journalism, and public advocacy helped shape Canadian life.')

paras = [
    ('When Mary Ann Shadd Cary arrived in Canada in 1851, she brought more than the experience of crossing a border. She brought a belief that education and public conversation could change what a community considered possible. Born in Wilmington, Delaware, in 1823, Shadd grew up in a family committed to abolition and learning. After the passage of the Fugitive Slave Act in the United States, she moved to Canada and settled in Windsor, Ontario.'),
    ('Shadd worked as an educator and opened a racially integrated school. Teaching was not simply employment for her; it was a way to give young people practical independence and a stronger voice in public life. She argued that Black communities should have access to serious education and should be able to build their futures without being confined by segregation or limited expectations.'),
    ('Her most enduring Canadian contribution was journalism. In 1853, she began publishing The Provincial Freeman, a newspaper that reported on Black life, challenged slavery, and encouraged Black people to pursue education, work, and self-determination. The paper gave readers information and argument at a time when Black voices were often described by others rather than heard directly. Shadd’s work also showed that newcomers and marginalized communities could create institutions of their own instead of waiting for permission to be included.'),
    ('Shadd later returned to the United States, where she supported the Union cause during the Civil War and continued her work for racial equality and women’s rights. She studied law and became one of the first Black women in the United States to do so. Her life therefore crosses several histories at once: migration, abolition, journalism, education, women’s suffrage, and the long effort to make citizenship meaningful.'),
    ('Her story matters to me because it changes the way I think about contribution. A person does not need to hold the most visible office or become famous nationwide to influence a country’s future. Shadd built a school, published a newspaper, and argued for a society in which people could direct their own lives. Those actions created knowledge and confidence that could travel farther than she could personally see.'),
    ('Mary Ann Shadd Cary should be remembered as a builder of Canadian public life. Her example asks a useful question of the present: whose work is making a community stronger even when it is not yet receiving broad recognition? Remembering her helps Canadians see that the country has always been shaped by people who crossed borders, created institutions, and insisted that education and a free press belong to everyone.'),
]
for text in paras:
    doc.add_paragraph(text)

doc.add_heading('Sources', level=1)
sources = [
    'Government of Canada, Mary Ann Shadd Cary (1823–1893): https://www.canada.ca/en/women-gender-equality/commemorations-celebrations/women-impact/human-rights/mary-ann-shadd-cary.html',
    'Dictionary of Canadian Biography, Mary Ann Camberton Shadd Cary: https://www.biographi.ca/en/bio/6425',
    'Library and Archives Canada, Mary Ann Shadd Cary collection: https://recherche-collection-search.bac-lac.gc.ca/eng/home/record?app=fonandcol&idnumber=100441',
]
for source in sources:
    p = doc.add_paragraph(source, style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        run.font.size = Pt(9)

doc.save(out)
print(out)
