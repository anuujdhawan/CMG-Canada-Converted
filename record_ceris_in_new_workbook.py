from openpyxl import load_workbook
p='/Users/themacintosh/Documents/My Files/Projects/CMG-Tailwind-components(V2)/outputs/backlink-created-details.xlsx'
wb=load_workbook(p); ws=wb['Backlink Details']
row=['https://ceris.ca/write-for-us/','https://commonwealthmigration.ca/immigrate/express-entry']
if row not in [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]: ws.append(row)
wb.save(p); print('rows',ws.max_row)
